import streamlit as st
import streamlit.components.v1 as components
import json
import re
import os
import hashlib
import urllib.parse
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from openai import OpenAI

# ═══════════════════════════════════════════════════
# PAGE CONFIG
# ═══════════════════════════════════════════════════
st.set_page_config(
    page_title="SanghaStatus",
    page_icon="🏛️",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# ═══════════════════════════════════════════════════
# SESSION STATE
# ═══════════════════════════════════════════════════
import sqlite3

MAX_HISTORY  = 10
HISTORY_FILE = ".sanghastatus_history.json"   # legacy file — auto-migrated into the DB below, then unused
DB_FILE      = ".sanghastatus.db"


def _get_db_conn():
    conn = sqlite3.connect(DB_FILE)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            date TEXT, project TEXT, members INTEGER, tasks INTEGER,
            tone TEXT, domain TEXT, lang TEXT,
            data_json TEXT, members_detail_json TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    return conn

# ═══════════════════════════════════════════════════
# UNIFIED SETTINGS — shared with SatiCast
# ═══════════════════════════════════════════════════
# Same shared-settings file SatiCast writes to (see its own comment for the
# full caveat): only actually shared across apps on the same filesystem —
# otherwise this quietly falls back to "remembers your name in this app only".
SHARED_SETTINGS_FILE = ".sati_sangha_shared_settings.json"


def load_shared_settings() -> dict:
    try:
        if os.path.exists(SHARED_SETTINGS_FILE):
            with open(SHARED_SETTINGS_FILE, "r") as f:
                return json.load(f)
    except Exception:
        pass
    return {}


def save_shared_settings(updates: dict) -> None:
    try:
        current = load_shared_settings()
        current.update(updates)
        with open(SHARED_SETTINGS_FILE, "w") as f:
            json.dump(current, f)
    except Exception:
        pass


ROSTER_FILE = ".sanghastatus_roster.json"


def load_roster() -> dict:
    """Persisted team-member profile memory: {name_lower: {"display": str,
    "role": str}}. Same shared-file caveat as history/settings — survives
    refreshes on this deployment, not a per-user database."""
    try:
        if os.path.exists(ROSTER_FILE):
            with open(ROSTER_FILE, "r") as f:
                return json.load(f)
    except Exception:
        pass
    return {}


def save_roster(roster: dict) -> None:
    try:
        with open(ROSTER_FILE, "w") as f:
            json.dump(roster, f)
    except Exception:
        pass


def load_history_from_disk() -> list:
    """
    Persistence across page refreshes within the SAME running deployment,
    now backed by a real SQLite database (.sanghastatus.db) instead of a
    flat JSON file — proper indexing/querying, and it transparently
    migrates any pre-existing JSON history file into the DB the first time
    it runs. IMPORTANT CAVEAT (unchanged from the JSON version): this is
    still a single database FILE on the server's local disk, so it is:
      - SHARED across every visitor to this deployment (not private per user)
      - LOST on a fresh redeploy or if the platform's filesystem is ephemeral
    True per-user persistence would need a real client-server database with
    user accounts. This is appropriate for a small private/team deployment,
    not a public multi-tenant one.
    """
    try:
        conn = _get_db_conn()
        # One-time migration from the legacy JSON file, only if the DB is
        # still empty — never overwrites DB rows that already exist.
        if conn.execute("SELECT COUNT(*) FROM history").fetchone()[0] == 0 and os.path.exists(HISTORY_FILE):
            try:
                with open(HISTORY_FILE, "r") as f:
                    old_entries = json.load(f)
                for e in reversed(old_entries):  # oldest first, so id order matches recency
                    conn.execute(
                        "INSERT INTO history (date, project, members, tasks, tone, domain, lang, data_json, members_detail_json) "
                        "VALUES (?,?,?,?,?,?,?,?,?)",
                        (e.get("date"), e.get("project"), e.get("members"), e.get("tasks"),
                         e.get("tone"), e.get("domain"), e.get("lang"),
                         json.dumps(e.get("data", {})), json.dumps(e.get("members_detail", [])))
                    )
                conn.commit()
            except Exception:
                pass
        rows = conn.execute(
            "SELECT date, project, members, tasks, tone, domain, lang, data_json, members_detail_json "
            "FROM history ORDER BY id DESC LIMIT ?", (MAX_HISTORY,)
        ).fetchall()
        conn.close()
        return [
            {
                "date": r[0], "project": r[1], "members": r[2], "tasks": r[3],
                "tone": r[4], "domain": r[5], "lang": r[6],
                "data": json.loads(r[7] or "{}"), "members_detail": json.loads(r[8] or "[]"),
            }
            for r in rows
        ]
    except Exception:
        return []


def save_history_to_disk(history: list) -> None:
    """Replaces the table contents with the current in-memory list — keeps
    this a drop-in replacement for the old 'dump the whole list' JSON
    writer so no call site elsewhere in the app needs to change."""
    try:
        conn = _get_db_conn()
        conn.execute("DELETE FROM history")
        for e in reversed(history):  # reversed so autoincrement id order matches recency
            conn.execute(
                "INSERT INTO history (date, project, members, tasks, tone, domain, lang, data_json, members_detail_json) "
                "VALUES (?,?,?,?,?,?,?,?,?)",
                (e.get("date"), e.get("project"), e.get("members"), e.get("tasks"),
                 e.get("tone"), e.get("domain"), e.get("lang"),
                 json.dumps(e.get("data", {})), json.dumps(e.get("members_detail", [])))
            )
        conn.commit()
        conn.close()
    except Exception:
        pass  # read-only filesystem or other issue — fail silently, session-state still works


def init_state():
    defaults = {
        "history":        load_history_from_disk(),
        "session_tokens": 0,
    }
    for k, v in defaults.items():
        if k not in st.session_state:
            st.session_state[k] = v

init_state()

# ═══════════════════════════════════════════════════
# API
# ═══════════════════════════════════════════════════
if "NVIDIA_API_KEY" not in st.secrets:
    st.error("🔑 NVIDIA_API_KEY not found in Streamlit secrets.")
    st.stop()

client = OpenAI(
    base_url="https://integrate.api.nvidia.com/v1",
    api_key=st.secrets["NVIDIA_API_KEY"]
)

# Optional: GitHub activity auto-pull. Only enabled when the deployer has
# added their own credentials to Streamlit secrets — same opt-in pattern
# as the weather/news API keys, so this stays fully invisible until
# someone deliberately configures it. Requires a GitHub personal access
# token (repo/read:user scope is enough) since the public API alone rate-
# limits unauthenticated requests too aggressively for daily use.
GITHUB_TOKEN    = st.secrets.get("GITHUB_TOKEN", "")
GITHUB_USERNAME = st.secrets.get("GITHUB_USERNAME", "")
GITHUB_PULL_ENABLED = bool(GITHUB_TOKEN and GITHUB_USERNAME)


def fetch_github_activity(since_hours: int = 24) -> list:
    """Best-effort pull of the configured user's recent commits across
    their repos, formatted as ready-to-paste bullet lines. Returns [] on
    any failure or if not configured — this is a convenience, never a
    hard dependency."""
    if not GITHUB_PULL_ENABLED:
        return []
    import requests as _requests
    from datetime import timedelta as _timedelta
    try:
        since_iso = (datetime.utcnow() - _timedelta(hours=since_hours)).isoformat() + "Z"
        headers = {"Authorization": f"token {GITHUB_TOKEN}", "Accept": "application/vnd.github+json"}
        events_url = f"https://api.github.com/users/{GITHUB_USERNAME}/events?per_page=30"
        r = _requests.get(events_url, headers=headers, timeout=8)
        if r.status_code != 200:
            return []
        bullets = []
        for ev in r.json():
            if ev.get("created_at", "") < since_iso:
                continue
            if ev.get("type") == "PushEvent":
                repo = ev.get("repo", {}).get("name", "")
                for c in ev.get("payload", {}).get("commits", [])[:3]:
                    msg = (c.get("message") or "").splitlines()[0][:100]
                    if msg:
                        bullets.append(f"- {msg} ({repo})")
            elif ev.get("type") == "PullRequestEvent":
                action = ev.get("payload", {}).get("action", "")
                pr = ev.get("payload", {}).get("pull_request", {})
                title = (pr.get("title") or "")[:100]
                if title:
                    bullets.append(f"- {action.capitalize()} PR: {title}")
        return bullets[:10]
    except Exception:
        return []

# ═══════════════════════════════════════════════════
# CONSTANTS
# ═══════════════════════════════════════════════════
TONE_OPTIONS = {
    "📝 Formal":      "Use formal, complete sentences with professional vocabulary.",
    "💬 Semi-formal": "Use clear, professional but conversational sentences.",
    "⚡ Concise":     "Use short, punchy one-line bullets. Be brief and direct.",
}
DOMAIN_OPTIONS = {
    "💻 Software Dev":        "software development",
    "🧪 QA / Testing":        "quality assurance and testing",
    "🚀 DevOps":              "DevOps and infrastructure",
    "🎨 Design / UX":         "UI/UX design",
    "📦 Product":             "product management",
    "💰 Finance":             "financial operations",
    "🧩 Functional Team":     "functional / business operations (process coordination, stakeholder requirements, cross-functional workflows, and operational execution)",
    "🧑‍🤝‍🧑 Human Resource Team": "Human Resources (recruitment, onboarding, employee relations, policy compliance, performance management, and engagement initiatives)",
}
OUTPUT_LANG_OPTIONS = {
    "🇬🇧 English": "English",
    "🇮🇳 हिन्दी":  "Hindi (Devanagari script)",
    "🇮🇳 मराठी":   "Marathi (Devanagari script)",
}

LOADER_STAGES = [
    ("🏛️", "Gathering the Sangha…",             "Assembling your team's raw updates"),
    ("✍️", "Rewriting into professional prose…", "Expanding notes into full sentences"),
    ("✨", "Finalising your update…",             "Running final quality check"),
]

# Short flavor lines that rotate under the main loader stage, pure CSS —
# see SatiCast's identical mechanism for the full explanation.
LOADER_MICRO_COPY = [
    "Untangling bullet points…",
    "Finding each person's throughline…",
    "Smoothing the rough edges…",
    "Checking nobody got lost…",
]

# ═══════════════════════════════════════════════════
# HELPERS
# ═══════════════════════════════════════════════════
def parse_members(raw: str) -> list:
    members = []
    blocks = re.split(r'\n(?=[A-Z][^\n:]{0,30}:)', raw.strip())
    for block in blocks:
        lines = block.strip().splitlines()
        if not lines:
            continue
        header = lines[0].strip().rstrip(":")
        tasks = [l.strip().lstrip("-•*").strip() for l in lines[1:]
                 if l.strip() and l.strip()[0] in "-•*"]
        # Instant, pre-API heuristic for the member-chip badge — deliberately
        # conservative (phrases, not bare words) to avoid flagging normal
        # work like "resolved the issue" or "fixed the bug" as a blocker.
        # The LLM-generated blocker_summary (build_narrative_prompt) does the
        # real semantic judgment call; this is just a cheap instant hint.
        blocker_phrases = [
            "blocked", "block on", "blocked by", "stuck on", "stuck with",
            "waiting on", "waiting for", "pending approval", "pending access",
            "need approval", "need access", "unable to", "not able to",
            "can't proceed", "cannot proceed", "dependency on", "depends on",
        ]
        blockers = [t for t in tasks if any(p in t.lower() for p in blocker_phrases)]
        if header and tasks:
            members.append({"name": header, "tasks": tasks, "blockers": blockers, "count": len(tasks)})
    return members


DAY_SEP_RE = re.compile(
    r'^\s*(?:[-=~_*]{2,}\s*)?'
    r'(?:(?:Mon|Tue|Wed|Thu|Fri|Sat|Sun)[a-zA-Z]*|Day\s*\d+|\d{1,2}[/-]\d{1,2}(?:[/-]\d{2,4})?)'
    r'\s*(?:[-=~_*]{2,})?\s*$',
    re.IGNORECASE
)


def split_multi_day_paste(raw: str):
    """Detects a bulk paste of several days' updates in one box (separated
    by a date/day-name line, e.g. 'Monday', 'Day 2', '15/09', '=== Tue ===')
    and splits it into (day_label, day_raw_text) segments. Returns an empty
    list when fewer than 2 segments are found, since a single day's paste
    should behave exactly as before."""
    lines = raw.splitlines()
    segments = []
    current_label = None
    current_lines = []
    for line in lines:
        stripped = line.strip()
        if stripped and DAY_SEP_RE.match(stripped):
            if current_lines and any(l.strip() for l in current_lines):
                segments.append((current_label or f"Day {len(segments)+1}", "\n".join(current_lines).strip()))
            current_label = stripped.strip("-=~_* ")
            current_lines = []
        else:
            current_lines.append(line)
    if current_lines and any(l.strip() for l in current_lines):
        segments.append((current_label or f"Day {len(segments)+1}", "\n".join(current_lines).strip()))
    return segments if len(segments) > 1 else []


DOMAIN_KEYWORDS = {
    "💻 Software Dev": ["deploy", "bug", "sprint", "merge", "pipeline", "api", "backend", "frontend", "code", "repo", "pr "],
    "🧪 QA / Testing": ["test case", "regression", "qa ", "test coverage", "bug report", "uat"],
    "🚀 DevOps":       ["infra", "kubernetes", "docker", "ci/cd", "server", "deployment", "monitoring"],
    "🎨 Design / UX":  ["wireframe", "prototype", "figma", "user test", "design review", "mockup"],
    "📦 Product":      ["roadmap", "backlog", "stakeholder", "feature spec", "user story"],
    "💰 Finance":      ["invoice", "budget", "reconcil", "forecast", "audit", "expense"],
    "🧩 Functional Team": ["coordinat", "escalat", "workflow", "operational", "process"],
    "🧑‍🤝‍🧑 Human Resource Team": ["onboard", "candidate", "interview", "recruit", "policy", "employee engagement", "grievance", "performance review"],
}


def suggest_domain(raw: str):
    """Best-effort domain guess from raw text keywords. Returns None if no clear signal."""
    if not raw.strip():
        return None
    text = raw.lower()
    scores = {d: sum(text.count(k) for k in kws) for d, kws in DOMAIN_KEYWORDS.items()}
    best = max(scores, key=scores.get)
    return best if scores[best] > 0 else None


def looks_malformed(raw: str, members: list):
    """Returns a warning message if the raw input looks off, else None."""
    if not raw.strip():
        return None
    if not members:
        return "Could not detect any team member names. Use a 'Name:' header followed by '-' bullet points."
    if len(raw.strip().splitlines()) <= 2:
        return "This looks very short — make sure each person's tasks are on their own '-' bulleted lines."
    return None


def word_diff_html(old_text: str, new_text: str) -> tuple:
    """Word-level diff between two task strings, returned as (old_html,
    new_html) with removed words wrapped in a red strike-through span in
    old_html and added words wrapped in a green span in new_html — so the
    'vs yesterday' panel shows what actually changed, not just two full
    strings side by side."""
    import html as _html
    from difflib import SequenceMatcher
    old_words = old_text.split()
    new_words = new_text.split()
    sm = SequenceMatcher(None, old_words, new_words)
    old_parts, new_parts = [], []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        old_chunk = _html.escape(" ".join(old_words[i1:i2]))
        new_chunk = _html.escape(" ".join(new_words[j1:j2]))
        if tag == "equal":
            old_parts.append(old_chunk)
            new_parts.append(new_chunk)
        else:
            if old_chunk:
                old_parts.append(f'<span class="diff-del">{old_chunk}</span>')
            if new_chunk:
                new_parts.append(f'<span class="diff-add">{new_chunk}</span>')
    return " ".join(old_parts), " ".join(new_parts)


def diff_against_previous(current_members: list, previous_entry: dict) -> dict:
    """
    For each person in current_members, finds tasks that closely match a
    task they had in the previous saved entry — signals possible carryover
    or a stalled task. Returns {name: [(today_task, yesterday_task), ...]}
    so the UI can show an actual before/after comparison, not just today's
    text in isolation.
    """
    if not previous_entry or not previous_entry.get("members_detail"):
        return {}
    from difflib import SequenceMatcher
    prev_by_name = {m["name"].lower(): m["tasks"] for m in previous_entry["members_detail"]}
    result = {}
    for m in current_members:
        prev_tasks = prev_by_name.get(m["name"].lower(), [])
        pairs = []
        for task in m["tasks"]:
            best_match, best_ratio = None, 0.0
            for prev_task in prev_tasks:
                ratio = SequenceMatcher(None, task.lower(), prev_task.lower()).ratio()
                if ratio > best_ratio:
                    best_match, best_ratio = prev_task, ratio
            if best_ratio > 0.6:
                pairs.append((task, best_match))
        if pairs:
            result[m["name"]] = pairs
    return result


def flag_brief_updates(current_members: list, history: list) -> dict:
    """
    Purely objective, quantitative comparison — NOT sentiment/emotion
    analysis. Flags when a person's update today has notably fewer words
    per task than their own historical average across saved entries. This
    is a factual verbosity observation only (e.g. "shorter than usual"),
    deliberately avoiding any inference about mood, stress, or wellbeing —
    a manager can use it as a cue to check in, not as a diagnosis.
    Returns {name: (today_avg_words, historical_avg_words)}.
    """
    if len(history) < 2:
        return {}  # need at least 2 past entries for a meaningful average

    word_counts_by_name = {}
    for entry in history:
        for m in entry.get("members_detail", []):
            name_key = m["name"].lower()
            if not m["tasks"]:
                continue
            avg_words = sum(len(t.split()) for t in m["tasks"]) / len(m["tasks"])
            word_counts_by_name.setdefault(name_key, []).append(avg_words)

    flagged = {}
    for m in current_members:
        key = m["name"].lower()
        past_avgs = word_counts_by_name.get(key, [])
        if len(past_avgs) < 2 or not m["tasks"]:
            continue
        historical_avg = sum(past_avgs) / len(past_avgs)
        today_avg = sum(len(t.split()) for t in m["tasks"]) / len(m["tasks"])
        if historical_avg > 4 and today_avg < historical_avg * 0.5:
            flagged[m["name"]] = (round(today_avg, 1), round(historical_avg, 1))
    return flagged


def dynamic_ta_height(text: str, min_h: int = 120, max_h: int = 500) -> int:
    """
    Sizes a text_area to fit its actual content instead of a fixed height
    that clips longer generations — counts wrapped lines (~90 chars/line)
    plus explicit newlines, at ~24px per line, clamped to a sane range.
    """
    if not text:
        return min_h
    line_count = 0
    for line in text.splitlines():
        line_count += max(1, -(-len(line) // 90))  # ceil division for wrapping
    return max(min_h, min(max_h, line_count * 24 + 40))


def to_jira_confluence_markup(chat_update: str) -> str:
    """
    Deterministic transform of the already-generated chat_update text into
    Jira/Confluence wiki markup — no extra LLM call needed, so it's instant.
    "• Name" -> "h3. Name" ; "- task" -> "* task"
    """
    out_lines = []
    for line in chat_update.splitlines():
        stripped = line.strip()
        if stripped.startswith("•"):
            out_lines.append(f"h3. {stripped.lstrip('•').strip()}")
        elif stripped.startswith("-"):
            out_lines.append(f"* {stripped.lstrip('-').strip()}")
        elif stripped:
            out_lines.append(f"h2. {stripped}")
        else:
            out_lines.append("")
    return "\n".join(out_lines)


def build_docx_export(data: dict, fmt_date: str, prepared_by: str = ""):
    """Builds a Word document from the generated status update. Returns
    (bytes, None) on success, or (None, error_message) if python-docx
    isn't installed in this environment."""
    try:
        from docx import Document
        import io as _io
    except ImportError:
        return None, "python-docx isn't installed — add 'python-docx' to requirements.txt to enable this."

    doc = Document()
    doc.add_heading(f"Daily Status Update — {fmt_date}", level=1)
    if prepared_by:
        p = doc.add_paragraph()
        p.add_run(f"Prepared by {prepared_by}").italic = True

    doc.add_heading("Standup Narrative", level=2)
    doc.add_paragraph(data.get("standup_narrative", ""))

    if data.get("tomorrow_plan"):
        doc.add_heading("Tomorrow's Plan", level=2)
        doc.add_paragraph(data["tomorrow_plan"])

    if data.get("blocker_summary"):
        doc.add_heading("Blockers", level=2)
        doc.add_paragraph(data["blocker_summary"])

    doc.add_heading("Chat Update", level=2)
    doc.add_paragraph(data.get("chat_update", ""))

    doc.add_heading("Email Update", level=2)
    doc.add_paragraph(data.get("email_update", ""))

    buf = _io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read(), None


def render_copy_button(text: str, unique_key: str, label: str = "📋 Copy"):
    """A small copy-to-clipboard button with a genuine micro-interaction —
    the label flips to a checkmark and briefly flashes green on success —
    instead of a plain browser toast. Copies the text as generated (not
    live edits made in the paired text_area, since reaching into that
    DOM reliably across Streamlit versions isn't worth the fragility)."""
    payload = json.dumps(text)
    components.html(f"""
    <div style="font-family:'Inter',sans-serif;">
        <button id="copyBtn_{unique_key}" onclick="doCopy_{unique_key}()"
            style="padding:5px 14px;border-radius:8px;border:1px solid #D97757;
            background:transparent;color:#D97757;cursor:pointer;font-size:0.78rem;
            font-weight:600;transition:all 0.2s ease;">{label}</button>
    </div>
    <script>
        function doCopy_{unique_key}() {{
            const btn = document.getElementById('copyBtn_{unique_key}');
            navigator.clipboard.writeText({payload}).then(() => {{
                btn.textContent = '✅ Copied!';
                btn.style.background = '#22c55e';
                btn.style.borderColor = '#22c55e';
                btn.style.color = '#fff';
                setTimeout(() => {{
                    btn.textContent = '{label}';
                    btn.style.background = 'transparent';
                    btn.style.borderColor = '#D97757';
                    btn.style.color = '#D97757';
                }}, 1400);
            }}).catch(() => {{ btn.textContent = '⚠️ Copy failed'; }});
        }}
    </script>
    """, height=42)


def mailto_link(subject: str, body: str) -> str:
    s = urllib.parse.quote(subject)
    b = urllib.parse.quote(body)
    return f"mailto:?subject={s}&body={b}"


def render_loader(stage_idx: int) -> str:
    total = len(LOADER_STAGES)
    emoji, title, sub = LOADER_STAGES[stage_idx]
    pct = int((stage_idx + 1) / total * 100)
    dots = "".join([
        f'<div class="ld {"ldone" if i < stage_idx else ("lactive" if i == stage_idx else "")}"></div>'
        for i in range(total)
    ])
    n_micro = len(LOADER_MICRO_COPY)
    slot_secs = 2.5
    loop_secs = n_micro * slot_secs
    micro_spans = "".join(
        f'<span class="loader-micro-item" style="animation-duration:{loop_secs}s;'
        f'animation-delay:-{i*slot_secs}s;">{line}</span>'
        for i, line in enumerate(LOADER_MICRO_COPY)
    )
    return (
        f'<div class="loader-wrap">'
        f'<div class="loader-emoji">{emoji}</div>'
        f'<div class="loader-title">{title}</div>'
        f'<div class="loader-sub">{sub} &nbsp;·&nbsp; {stage_idx+1}/{total}</div>'
        f'<div class="loader-micro">{micro_spans}</div>'
        f'<div class="loader-dots">{dots}</div>'
        f'<div class="loader-bar-bg"><div class="loader-bar-fg" style="width:{pct}%"></div></div>'
        f'<div class="loader-pct">{pct}%</div>'
        f'</div>'
    )


# ═══════════════════════════════════════════════════
# SYSTEM PROMPT
# ═══════════════════════════════════════════════════
def extract_json_object(raw: str) -> dict:
    """
    Robustly parse a JSON object from an LLM response. Some models
    (e.g. reasoning models like Nemotron 3.5 Lightning) prepend a
    <think>...</think> block before the actual JSON, or wrap the JSON
    in markdown code fences — both of which break a plain json.loads.
    This strips those, then extracts the {...} substring before parsing.
    """
    s = raw.strip()

    # Strip a <think>...</think> block if present (even if unterminated —
    # in that case drop everything up to the last </think> or, if there's
    # no closing tag at all, take everything after the LAST '}' as a
    # fallback isn't reliable, so just remove the tag pair when found).
    s = re.sub(r'<think>.*?</think>', '', s, flags=re.DOTALL)
    s = re.sub(r'^.*?</think>', '', s, count=1, flags=re.DOTALL)  # unterminated opening tag case

    # Strip markdown code fences (```json ... ``` or ``` ... ```)
    s = re.sub(r'^```(?:json)?\s*', '', s.strip())
    s = re.sub(r'```\s*$', '', s.strip())

    s = s.strip()

    # Extract the outermost {...} in case there's any leftover stray text
    start = s.find('{')
    end = s.rfind('}')
    if start != -1 and end != -1 and end > start:
        s = s[start:end + 1]

    return json.loads(s)


def call_llm_json(system_prompt: str, user_prompt: str, temperature: float, max_tokens: int):
    """
    Calls the LLM and returns the parsed JSON dict. Tries to disable the
    model's reasoning/thinking mode via extra_body first (Nemotron 3.5
    Lightning is a hybrid reasoning model that can otherwise prepend a
    <think>...</think> block before the JSON); if the server rejects that
    parameter, retries once without it. extract_json_object() strips any
    leftover <think> block regardless, as a safety net either way.
    """
    kwargs = dict(
        model="nvidia/nemotron-3.5-lightning-30b-a3b",
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt}
        ],
        temperature=temperature,
        max_tokens=max_tokens,
        response_format={"type": "json_object"},
    )
    try:
        completion = client.chat.completions.create(
            **kwargs, extra_body={"chat_template_kwargs": {"thinking": False}}
        )
    except Exception:
        completion = client.chat.completions.create(**kwargs)

    data = extract_json_object(completion.choices[0].message.content)
    return data, completion


def _shared_instructions(tone: str, domain: str, lang: str) -> str:
    """Common domain/tone/language framing reused by all three split prompts."""
    return f"""
You are SanghaStatus, a precise professional assistant for workplace teams.

DOMAIN CONTEXT: You are writing updates for a {domain} team.
Use vocabulary, phrasing, and framing natural to this domain. For example:
- Software Dev / QA / DevOps: use terms like "deployed", "resolved", "tested", "merged", "pipeline", "sprint".
- Functional Team: use terms like "coordinated", "aligned with stakeholders", "processed", "escalated", "streamlined workflow", "operational milestone".
- Human Resource Team: use terms like "onboarded", "screened candidates", "conducted interviews", "policy rollout", "employee engagement", "performance review", "grievance resolved" — and NEVER use software/technical jargon (no "deployed", "bug", "sprint", etc.) unless the raw input explicitly mentions it.
- Finance: use terms like "reconciled", "audited", "processed invoices", "forecasted", "budget review".
- Design/UX: use terms like "prototyped", "wireframed", "user tested", "iterated on design".
Always stay faithful to what the raw input actually says — domain vocabulary should make the phrasing natural, not invent activities that didn't happen.

TONE: {tone}
OUTPUT LANGUAGE: Write ALL output fields in {lang}.

STRICT RULES:
- TRANSFORM RAW INPUTS: DO NOT just copy-paste the raw bullet points. You MUST rewrite and expand every single short phrase into a full, professional, action-oriented sentence ending with a period.
  * Bad: "- Support for testing activities"
  * Good: "- Provided support for testing activities, ensuring test coverage stayed on track."
  * Bad: "- working on masking"
  * Good: "- Currently working on implementing masking enhancements."
- Every single bullet — including the LAST one in a person's list — must be fully rewritten this way. Do not leave any bullet as a bare noun phrase or fragment from the raw input; every bullet must start with a past- or present-tense action verb.
- Each raw task stays as its own separate bullet. Do NOT merge tasks or summarize multiple tasks into one.
- Keep each person's updates strictly separated under their name.
- Use correct pronouns inferred from context, or use their name if unclear.
- NO introductory or concluding conversational text in the JSON values.
- Return ONLY a valid JSON object — no markdown fences, no preamble.
- Do NOT include any reasoning, chain-of-thought, or <think> tags of any kind in your response — output ONLY the raw JSON object and nothing else, starting with {{ and ending with }}.
- Every sentence must be grammatically complete — never truncate or cut a sentence short partway through.
"""


def build_narrative_prompt(tone, domain, lang, include_tomorrow, include_blockers):
    """Narrative-family output: the spoken standup, tomorrow's plan, blockers.
    Split out so it can run in parallel with the chat/email prompts below."""
    tomorrow_key = (
        ',\n  "tomorrow_plan": "Good morning, everyone.\\n\\n[Name 1] plans to [tomorrow task].\\n[Name 2] plans to [tomorrow task].\\n\\nThat concludes the plan for tomorrow."'
        if include_tomorrow else ""
    )
    blocker_key = (
        ',\n  "blocker_summary": "Blockers identified today:\\n\\n[Name]: [Blocker detail].\\n(Write \\"No blockers reported today.\\" if none)"'
        if include_blockers else ""
    )
    blocker_guidance = (
        """
BLOCKER DETECTION: Identify blockers by MEANING, not by keyword-matching. A task
is a blocker if it describes being stuck, waiting, or unable to proceed —
including phrasing that doesn't contain an obvious keyword like "block", e.g.:
- "Waiting on the vendor to confirm pricing" → blocker (waiting on external party)
- "Need approval from finance before proceeding" → blocker (needs approval)
- "Access request still pending with IT" → blocker (pending access)
- "Can't reproduce the issue without prod data" → blocker (missing dependency)
Do NOT treat a task as a blocker just because it mentions a bug or issue being
WORKED ON or RESOLVED — only flag it if the person is genuinely stuck/waiting.
"""
        if include_blockers else ""
    )
    return _shared_instructions(tone, domain, lang) + blocker_guidance + f"""
Return a JSON object with EXACTLY these keys:
{{
  "standup_narrative": "Good morning, everyone.\\n\\nHere is the status update for [Date].\\n\\n[Name 1]: [narrative].\\n\\n[Name 2]: [narrative].\\n\\nThat concludes today's status update."{tomorrow_key}{blocker_key}
}}
"""


def build_chat_prompt(tone, domain, lang):
    """Chat-family output: Slack/Teams update + WhatsApp update.
    Runs in parallel with the narrative and email prompts."""
    return _shared_instructions(tone, domain, lang) + """
Use hyphen "-" for bullets. Never asterisks in chat_update (WhatsApp update may use *bold* for names).
IMPORTANT: each person's name/header appears ONLY ONCE, with ALL of their tasks listed as bullets underneath it — never repeat a person's name header for each separate task.
REMINDER — REWRITE, DO NOT COPY: every bullet must be your own rewritten, full sentence — never the raw input text with only capitalization/punctuation fixed.

Return a JSON object with EXACTLY these keys:
{
  "chat_update": "Daily Project Status Update | [Date]\\n\\n• [Name 1]\\n- [Task 1 sentence].\\n- [Task 2 sentence].\\n\\n• [Name 2]\\n- [Task 1 sentence].",
  "whatsapp_update": "📋 *Daily Status Update | [Date]*\\n\\n*[Name 1]*\\n- [Task 1 sentence].\\n- [Task 2 sentence].\\n\\n*[Name 2]*\\n- [Task 1 sentence]."
}
"""


def build_email_prompt(tone, domain, lang):
    """Email-family output. Runs in parallel with the narrative and chat prompts."""
    return _shared_instructions(tone, domain, lang) + """
Use hyphen "-" for bullets. Never asterisks.

REMINDER — REWRITE, DO NOT COPY: The bullets under [Task 1], [Task 2] etc. must be your OWN
rewritten, full, professional sentences — never the raw input text verbatim, even with minor
punctuation fixes. If a raw line reads "-Worked on the Data migration activity", the email
bullet must be substantively rewritten (e.g. "- Led the data migration effort, ensuring a
smooth transition of records.") — not simply recapitalized and given a period. Every bullet in
the email must sound like it came from a professional writer, not a copy-paste of the input.

SIGN-OFF RULE: The email is being SENT ABOUT the team members listed, not BY one of them —
never sign the closing with any of the team members' own names, since that would make it
look like a team member is reporting on themselves to the team. Always close with a generic
sign-off that names no individual, exactly: "Regards,\\nProject Team" — do not substitute any
person's name here under any circumstances, even if only one team member was mentioned.

Return a JSON object with EXACTLY this key:
{
  "email_update": "Subject: Daily Project Status Update | [Date]\\n\\nDear Team,\\n\\nPlease find below the status update for [Date].\\n\\n[Name 1]\\n- [Fully rewritten task sentence, not a copy of the input].\\n\\n[Name 2]\\n- [Fully rewritten task sentence].\\n\\nKindly revert in case of any queries.\\n\\nRegards,\\nProject Team"
}
"""


# ═══════════════════════════════════════════════════
# CSS — pure-CSS dark toggle via :has(), zero Streamlit rerun
# ═══════════════════════════════════════════════════
CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800;900&family=Syne:wght@700;800&display=swap');

:root {
    --app-bg:linear-gradient(-45deg,#F5EFE3,#FBEEE5,#F0E4D4,#F7F2E8);
    --card-bg:rgba(255,255,255,0.96); --card-bdr:rgba(255,255,255,1);
    --text-h:#1F1E1D; --text-b:#1e293b; --text-m:#475569;
    --input-bg:#f8fafc; --input-bdr:#cbd5e1; --input-txt:#0f172a;
    --btn-sh1:rgba(180,85,50,0.4);
    --out-bg:#ffffff; --code-bg:#fafafa; --code-txt:#0f172a;
    --chip-bg:rgba(255,255,255,0.7); --chip-bdr:rgba(0,0,0,0.08);
    --loader-bg:rgba(255,255,255,0.97); --loader-title:#1e1b4b; --loader-sub:#6d28d9;
}
body:has(#dmchk:checked) {
    --app-bg:linear-gradient(-45deg,#1F1E1D,#211E1A,#241C18,#1D1B17);
    --card-bg:rgba(15,23,42,0.92); --card-bdr:rgba(255,255,255,0.08);
    --text-h:#F1F5F9; --text-b:#CBD5E1; --text-m:#94A3B8;
    --input-bg:rgba(15,23,42,0.95); --input-bdr:rgba(217,119,87,0.35); --input-txt:#F1F5F9;
    --btn-sh1:rgba(217,119,87,0.35);
    --out-bg:rgba(15,23,42,0.9); --code-bg:#0f172a; --code-txt:#E2E8F0;
    --chip-bg:rgba(255,255,255,0.06); --chip-bdr:rgba(255,255,255,0.1);
    --loader-bg:rgba(10,15,30,0.96); --loader-title:#DDD6FE; --loader-sub:#A78BFA;
}

html, body, [class*="css"] { font-family:'Inter',sans-serif !important; color:var(--text-b) !important; }
.stApp {
    background:var(--app-bg) !important; background-size:400% 400% !important;
    animation:gradientBG 18s ease infinite !important; min-height:100vh;
}
@keyframes gradientBG { 0%{background-position:0% 50%} 50%{background-position:100% 50%} 100%{background-position:0% 50%} }
.main .block-container { max-width:1240px; padding-top:0.5rem !important; padding-bottom:4rem; padding-left:2rem; padding-right:2rem; }
#MainMenu, footer, header { visibility:hidden; }
* { box-sizing:border-box; }

/* ── PURE-CSS DARK TOGGLE — zero rerun, generation never interrupted ── */
#dmchk { display:none; }
.dm-label {
    position:fixed; top:1.1rem; right:1.4rem; z-index:99999;
    width:58px; height:30px; border-radius:999px; cursor:pointer;
    background:rgba(255,255,255,0.22); border:1.5px solid rgba(255,255,255,0.4);
    display:flex; align-items:center; padding:3px; backdrop-filter:blur(10px);
    box-shadow:0 4px 14px rgba(0,0,0,0.15); transition:all 0.3s ease;
}
.dm-label::after {
    content:'🌙'; width:24px; height:24px; border-radius:50%;
    background:#fff; display:flex; align-items:center; justify-content:center;
    font-size:0.8rem; transition:transform 0.35s cubic-bezier(0.34,1.56,0.64,1);
}
body:has(#dmchk:checked) .dm-label::after { content:'☀️'; transform:translateX(28px); }

/* ── FONT-SIZE ACCESSIBILITY CONTROL ── */
html { font-size:16px; transition:font-size 0.2s ease; }
html:has(#fsSmall:checked) { font-size:14px; }
html:has(#fsLarge:checked) { font-size:18px; }
#fsSmall, #fsNormal, #fsLarge { display:none; }
.fs-toggle {
    position:fixed; top:1.1rem; right:5.2rem; z-index:99999;
    display:flex; gap:2px; background:var(--card-bg);
    border:1.5px solid var(--card-bdr); border-radius:999px; padding:3px;
    backdrop-filter:blur(10px); box-shadow:0 4px 14px rgba(0,0,0,0.1);
}
.fs-btn {
    width:26px; height:24px; display:flex; align-items:center; justify-content:center;
    border-radius:999px; font-size:0.7rem; font-weight:800; cursor:pointer;
    color:var(--text-m); transition:all 0.2s ease;
}
#fsSmall:checked ~ label[for="fsSmall"],
#fsNormal:checked ~ label[for="fsNormal"],
#fsLarge:checked ~ label[for="fsLarge"] {
    background:#D97757; color:#fff;
}

/* ── HERO ── */
.hero-section { text-align:center; margin-bottom:2.5rem; position:relative; z-index:1; }
.hero-icon {
    font-size:4.2rem; display:block; margin-bottom:0.7rem;
    animation:floatIcon 3.2s ease-in-out infinite;
    filter:drop-shadow(0 8px 20px rgba(0,0,0,0.2));
}
@keyframes floatIcon { 0%,100%{transform:translateY(0) rotate(-2deg)} 50%{transform:translateY(-10px) rotate(2deg)} }
.main-title {
    font-family:'Syne',sans-serif; font-size:4.4rem; font-weight:900;
    letter-spacing:-2px; margin-bottom:0.4rem;
    background:linear-gradient(120deg,#D97757,#8B3A1F,#D97757,#B45532);
    background-size:280% auto;
    -webkit-background-clip:text; -webkit-text-fill-color:transparent; background-clip:text;
    animation:gradShift 6s linear infinite;
}
@keyframes gradShift { to { background-position:280% center; } }
.main-subtitle {
    color:#5A4632 !important; font-size:1.2rem; font-weight:600;
    max-width:680px; margin:0.3rem auto 0; line-height:1.65;
}
.sangha-meaning {
    display:inline-flex; align-items:center; gap:8px; margin-top:1rem;
    background:rgba(217,119,87,0.12); border:1.5px solid rgba(217,119,87,0.35);
    border-radius:999px; padding:0.5rem 1.5rem; font-size:0.9rem; font-weight:600;
    color:#8B3A1F !important; -webkit-text-fill-color:#8B3A1F !important;
    backdrop-filter:blur(12px);
    animation:fadeUp 0.7s 0.15s ease both;
}
@keyframes fadeUp { from{opacity:0;transform:translateY(12px)} to{opacity:1;transform:translateY(0)} }

/* ── CARDS ── */
.custom-card {
    background:var(--card-bg); border-radius:26px; padding:1.9rem 2.4rem;
    box-shadow:0 18px 46px rgba(0,0,0,0.16); margin-bottom:1.5rem;
    border:1px solid var(--card-bdr); backdrop-filter:blur(20px);
    transition:transform 0.3s ease, box-shadow 0.3s ease;
    animation:cardIn 0.5s ease both;
}
.custom-card:hover { transform:translateY(-3px); box-shadow:0 24px 56px rgba(0,0,0,0.2); }
@keyframes cardIn { from{opacity:0;transform:translateY(20px)} to{opacity:1;transform:translateY(0)} }
.card-header { display:flex; justify-content:space-between; align-items:center; margin-bottom:1.2rem; }
.card-title  { display:flex; align-items:center; gap:16px; }
.icon-circle {
    width:56px; height:56px; border-radius:16px; display:flex; align-items:center; justify-content:center;
    font-size:1.6rem; box-shadow:0 8px 18px rgba(0,0,0,0.14); flex-shrink:0;
    transition:transform 0.3s ease;
}
.custom-card:hover .icon-circle { transform:rotate(-6deg) scale(1.06); }
.blue-icon   { background:linear-gradient(135deg,#3b82f6,#2563eb); }
.green-icon  { background:linear-gradient(135deg,#10b981,#059669); }
.purple-icon { background:linear-gradient(135deg,#8b5cf6,#7c3aed); }
.orange-icon { background:linear-gradient(135deg,#f59e0b,#d97706); }
.title-text { font-size:1.42rem; font-weight:800; color:var(--text-h); }
.desc-text  { color:var(--text-m); font-size:0.92rem; margin-top:0.2rem; }
.side-emoji { font-size:44px; opacity:0.9; }

/* ── INPUTS ── */
.stTextArea textarea, .stTextInput input {
    background-color:var(--input-bg) !important; border:2px solid var(--input-bdr) !important;
    border-radius:16px !important; color:var(--input-txt) !important;
    font-size:1rem !important; font-weight:500 !important; transition:all 0.25s ease !important;
}
.stTextArea textarea { min-height:230px !important; line-height:1.8 !important; padding:1rem !important; }
.stTextArea textarea::placeholder, .stTextInput input::placeholder {
    color:var(--text-m) !important;
    opacity:0.75 !important;
}
.stTextArea textarea:focus, .stTextInput input:focus {
    border-color:#D97757 !important; box-shadow:0 0 0 4px rgba(217,119,87,0.15) !important;
}
/* Box background/border kept exactly as before — only fixing text color so it's
   readable: light box (light theme) gets dark text, dark box (dark theme) gets light text. */
div[data-testid="stSelectbox"] > div > div, div[data-testid="stDateInput"] > div > div {
    background:var(--input-bg) !important; border:2px solid var(--input-bdr) !important;
    border-radius:16px !important;
}
/* Comprehensive text-color override — targets every possible nested node BaseWeb
   might render the value/placeholder text in. */
div[data-testid="stSelectbox"] div[data-baseweb="select"],
div[data-testid="stSelectbox"] div[data-baseweb="select"] div,
div[data-testid="stSelectbox"] div[data-baseweb="select"] span,
div[data-testid="stSelectbox"] div[data-baseweb="select"] input,
div[data-testid="stSelectbox"] div[data-baseweb="select"] p,
div[data-testid="stDateInput"] input {
    color:var(--input-txt) !important;
    -webkit-text-fill-color:var(--input-txt) !important;
}
div[data-testid="stSelectbox"] div[aria-disabled="true"],
div[data-testid="stSelectbox"] div[aria-disabled="true"] * {
    color:var(--input-txt) !important;
    opacity:0.9 !important;
    -webkit-text-fill-color:var(--input-txt) !important;
}
/* The dropdown OPTIONS LIST is rendered in a portal appended to <body>, so it must be
   themed globally, matched to the same theme colors. */
div[data-baseweb="popover"] ul,
div[data-baseweb="menu"] {
    background:var(--input-bg) !important;
    border:1.5px solid var(--input-bdr) !important;
    border-radius:12px !important;
}
div[data-baseweb="popover"] li,
div[data-baseweb="menu"] li,
div[data-baseweb="popover"] li *,
div[data-baseweb="menu"] li *,
div[data-baseweb="popover"] [role="option"],
div[data-baseweb="menu"] [role="option"],
div[data-baseweb="popover"] [role="option"] *,
div[data-baseweb="menu"] [role="option"] * {
    color:var(--input-txt) !important;
    -webkit-text-fill-color:var(--input-txt) !important;
    background:transparent !important;
}
div[data-baseweb="popover"] li:hover,
div[data-baseweb="menu"] li:hover,
div[data-baseweb="popover"] [role="option"]:hover,
div[data-baseweb="menu"] [role="option"]:hover {
    background:rgba(128,128,128,0.2) !important;
}
div[data-testid="stSelectbox"] label, div[data-testid="stTextInput"] label,
div[data-testid="stDateInput"] label, div[data-testid="stTextArea"] label {
    color:var(--text-h) !important; font-weight:700 !important; font-size:0.8rem !important;
    letter-spacing:0.06em !important; text-transform:uppercase !important;
}
.stCheckbox label p { color:var(--text-h) !important; font-weight:600 !important; }

/* ── MEMBER CHIPS ── */
/* Kanban-style board layout — uniform grid columns instead of a loose
   flex-wrap row, so member cards line up like board cards rather than
   free-floating chips. */
.member-chips { display:grid; grid-template-columns:repeat(auto-fill, minmax(220px, 1fr)); gap:12px; margin-top:1.1rem; }

/* ── SIDE-BY-SIDE DIFF PANEL — actual before/after comparison ── */
.diff-panel {
    margin-top:1rem; background:rgba(245,158,11,0.06);
    border:1.5px solid rgba(245,158,11,0.25); border-radius:14px; padding:1rem 1.2rem;
}
.diff-panel-title { font-weight:800; font-size:0.85rem; color:#92400E; margin-bottom:0.7rem; }
.diff-row { margin-bottom:0.7rem; }
.diff-row:last-child { margin-bottom:0; }
.diff-name { font-size:0.78rem; font-weight:700; color:var(--text-m); margin-bottom:0.3rem; }
.diff-cols { display:flex; align-items:center; gap:10px; flex-wrap:wrap; }
.diff-col {
    flex:1; min-width:200px; font-size:0.82rem; line-height:1.4;
    border-radius:8px; padding:0.5rem 0.7rem; position:relative;
}
.diff-label {
    display:block; font-size:0.65rem; font-weight:800; letter-spacing:0.06em;
    text-transform:uppercase; margin-bottom:0.2rem; opacity:0.7;
}
.diff-yesterday { background:rgba(148,163,184,0.12); color:var(--text-m); }
.diff-today { background:rgba(217,119,87,0.12); color:var(--text-h); }
.diff-arrow { font-size:1.1rem; color:#D97757; font-weight:700; flex-shrink:0; }
@media (max-width:600px) { .diff-arrow { transform:rotate(90deg); } }
/* Word-level diff highlighting inside the yesterday/today panel above. */
.diff-del { background:rgba(239,68,68,0.18); color:#991B1B; text-decoration:line-through; border-radius:3px; padding:0 2px; }
.diff-add { background:rgba(34,197,94,0.18); color:#166534; border-radius:3px; padding:0 2px; font-weight:700; }
.member-chip {
    background:var(--chip-bg); border:1.5px solid var(--chip-bdr); border-radius:14px;
    padding:0.7rem 1.1rem; backdrop-filter:blur(8px);
    transition:transform 0.2s ease, box-shadow 0.2s ease;
    animation:chipIn 0.4s ease both;
    position:relative;
}
.chip-count-badge {
    position:absolute; top:-8px; right:-8px; min-width:22px; height:22px; border-radius:50%;
    background:#D97757; color:#fff; font-size:0.68rem; font-weight:800;
    display:flex; align-items:center; justify-content:center; padding:0 4px;
    box-shadow:0 2px 6px rgba(0,0,0,0.2);
}
.member-chip:hover { transform:translateY(-3px); box-shadow:0 10px 24px rgba(0,0,0,0.12); }
@keyframes chipIn { from{opacity:0;transform:scale(0.9)} to{opacity:1;transform:scale(1)} }
.member-chip:nth-child(1){animation-delay:0.00s} .member-chip:nth-child(2){animation-delay:0.08s}
.member-chip:nth-child(3){animation-delay:0.16s} .member-chip:nth-child(4){animation-delay:0.24s}
.member-chip:nth-child(5){animation-delay:0.32s} .member-chip:nth-child(6){animation-delay:0.40s}
.member-chip:nth-child(7){animation-delay:0.48s} .member-chip:nth-child(8){animation-delay:0.56s}
.chip-name  { font-weight:800; font-size:0.94rem; color:var(--text-h); }
.chip-tasks { font-size:0.76rem; color:var(--text-m); margin-top:0.15rem; }
.chip-block { font-size:0.72rem; color:#ef4444; font-weight:700; margin-top:0.2rem; }

/* ── MAIN BUTTON ── */
.stButton > button {
    background:linear-gradient(135deg,#D97757 0%,#B45532 100%) !important;
    color:#fff !important; border:none !important; border-radius:50px !important;
    padding:1.05rem 4.2rem !important; font-size:1.15rem !important; font-weight:800 !important;
    letter-spacing:0.05em; display:block; margin:2rem auto;
    box-shadow:0 10px 30px var(--btn-sh1) !important;
    transition:all 0.3s cubic-bezier(0.4,0,0.2,1) !important;
}
.stButton > button:hover { transform:translateY(-4px) scale(1.02) !important; box-shadow:0 18px 42px rgba(217,119,87,0.5) !important; }
.stButton > button:active { transform:scale(0.97) !important; }

/* ── SELECTBOX TEXT COLOR — dark mode only.
   Box background/position kept exactly as before; in dark mode the
   closed-value text, placeholder, and dropdown option list are all
   forced to white so they're readable against the dark box. ── */
body:has(#dmchk:checked) div[data-testid="stSelectbox"] > div > div,
body:has(#dmchk:checked) div[data-testid="stSelectbox"] div[data-baseweb="select"],
body:has(#dmchk:checked) div[data-testid="stSelectbox"] div[data-baseweb="select"] div,
body:has(#dmchk:checked) div[data-testid="stSelectbox"] div[data-baseweb="select"] span,
body:has(#dmchk:checked) div[data-testid="stSelectbox"] div[data-baseweb="select"] input,
body:has(#dmchk:checked) div[data-testid="stSelectbox"] div[data-baseweb="select"] p {
    color:#FFFFFF !important;
    -webkit-text-fill-color:#FFFFFF !important;
}
body:has(#dmchk:checked) div[data-baseweb="popover"] ul,
body:has(#dmchk:checked) div[data-baseweb="menu"] {
    background:#1E1B2E !important;
}
body:has(#dmchk:checked) div[data-baseweb="popover"] li,
body:has(#dmchk:checked) div[data-baseweb="menu"] li,
body:has(#dmchk:checked) div[data-baseweb="popover"] li *,
body:has(#dmchk:checked) div[data-baseweb="menu"] li *,
body:has(#dmchk:checked) div[data-baseweb="popover"] [role="option"],
body:has(#dmchk:checked) div[data-baseweb="menu"] [role="option"],
body:has(#dmchk:checked) div[data-baseweb="popover"] [role="option"] *,
body:has(#dmchk:checked) div[data-baseweb="menu"] [role="option"] * {
    color:#FFFFFF !important;
    -webkit-text-fill-color:#FFFFFF !important;
}

/* ── LOADER ── */
@keyframes loaderFadeIn { from{opacity:0;transform:translateY(24px)} to{opacity:1;transform:translateY(0)} }
@keyframes loaderBounce { 0%,100%{transform:translateY(0) scale(1)} 40%{transform:translateY(-14px) scale(1.14)} 65%{transform:translateY(-6px) scale(1.06)} }
@keyframes shimmer { 0%{background-position:200% 0} 100%{background-position:-200% 0} }
@keyframes dotPop { from{transform:scale(0.3);opacity:0} to{transform:scale(1);opacity:1} }
.loader-wrap {
    background:var(--loader-bg); border-radius:28px; padding:3rem 2.5rem; text-align:center;
    max-width:580px; margin:1.5rem auto; box-shadow:0 28px 70px rgba(0,0,0,0.25);
    animation:loaderFadeIn 0.4s cubic-bezier(0.34,1.56,0.64,1) both;
    border:1.5px solid rgba(255,255,255,0.12);
    position:relative; overflow:hidden;
}
.loader-wrap::after {
    content:''; position:absolute; top:0; left:-150%; width:100%; height:100%;
    background:linear-gradient(100deg, transparent, rgba(217,119,87,0.1), transparent);
    animation:loaderSweep 2.2s ease-in-out infinite;
}
@keyframes loaderSweep { to { left:150%; } }
.loader-emoji { font-size:3.4rem; display:block; margin-bottom:0.9rem; animation:loaderBounce 1.4s ease-in-out infinite; }
.loader-title { font-family:'Syne',sans-serif; font-size:1.35rem; font-weight:800; margin-bottom:0.35rem; color:var(--loader-title); }
.loader-sub   { font-size:0.84rem; font-weight:600; margin-bottom:1.6rem; color:var(--loader-sub); }
.loader-dots  { display:flex; justify-content:center; gap:9px; margin-bottom:1.5rem; }
.ld { width:11px; height:11px; border-radius:50%; background:#e5e7eb; transition:background 0.3s; }
.lactive { background:linear-gradient(135deg,#D97757,#B45532); animation:dotPop 0.4s cubic-bezier(0.34,1.56,0.64,1) both; box-shadow:0 0 10px rgba(217,119,87,0.5); }
.ldone { background:#B45532; }
.loader-bar-bg { height:7px; border-radius:99px; background:rgba(180,85,50,0.1); overflow:hidden; margin-bottom:0.6rem; }
.loader-bar-fg {
    height:100%; border-radius:99px; background:linear-gradient(90deg,#D97757,#B45532,#23a6d5,#D97757);
    background-size:200% 100%; animation:shimmer 1.5s linear infinite;
    transition:width 0.4s cubic-bezier(0.4,0,0.2,1);
}
.loader-pct { font-size:0.8rem; font-weight:800; letter-spacing:0.06em; color:var(--loader-sub); }

/* ── ROTATING MICRO-COPY — same pure-CSS technique as SatiCast. ── */
.loader-micro { position:relative; height:1.2rem; margin:0.4rem 0 0.8rem; }
.loader-micro-item {
    position:absolute; left:0; right:0; text-align:center;
    font-size:0.78rem; font-style:italic; color:#B45532; opacity:0;
    animation-name:microCycle; animation-timing-function:ease-in-out; animation-iteration-count:infinite;
}
body:has(#dmchk:checked) .loader-micro-item { color:#E8A87C; }
@keyframes microCycle {
    0% { opacity:0; }
    3% { opacity:1; }
    16% { opacity:1; }
    20% { opacity:0; }
    100% { opacity:0; }
}

/* ── OUTPUT GRID: 2 columns wide-screen ── */
.output-grid { display:grid; grid-template-columns:1fr 1fr; gap:1.5rem; }
@media (max-width:900px) { .output-grid { grid-template-columns:1fr; } }

@keyframes blockReveal { from{opacity:0;transform:translateY(22px)} to{opacity:1;transform:translateY(0)} }
.colored-block {
    padding:1.75rem 2rem; border-radius:18px 22px 22px 4px;
    background:var(--out-bg);
    box-shadow:0 4px 0 rgba(0,0,0,0.03), 0 10px 32px rgba(0,0,0,0.09);
    animation:blockReveal 0.5s ease both;
    transition:transform 0.25s ease, box-shadow 0.25s ease;
    position:relative; overflow:hidden;
}
.colored-block:hover { transform:translateY(-4px); box-shadow:0 6px 0 rgba(0,0,0,0.04), 0 18px 40px rgba(0,0,0,0.14); }
/* Paper-fold corner detail — small triangular fold, top-right */
.colored-block::after {
    content:''; position:absolute; top:0; right:0;
    width:0; height:0;
    border-style:solid; border-width:0 22px 22px 0;
    border-color:transparent rgba(0,0,0,0.05) transparent transparent;
    transition:border-width 0.2s ease;
}
.colored-block:hover::after { border-width:0 28px 28px 0; }

.narrative-block { border-left:7px solid #4f46e5; background:rgba(79,70,229,0.045); animation-delay:0.05s; grid-column:1/-1; }
.chat-block      { border-left:7px solid #ec4899; background:rgba(236,72,153,0.045); animation-delay:0.12s; }
.whatsapp-block  { border-left:7px solid #25D366; background:rgba(37,211,102,0.045); animation-delay:0.19s; }
.email-block     { border-left:7px solid #f59e0b; background:rgba(245,158,11,0.045); animation-delay:0.26s; grid-column:1/-1; }
.tomorrow-block  { border-left:7px solid #06b6d4; background:rgba(6,182,212,0.045);  animation-delay:0.33s; }
.blocker-block   { border-left:7px solid #ef4444; background:rgba(239,68,68,0.045);  animation-delay:0.40s; }

.block-title {
    font-size:1.28rem; font-weight:900; margin-bottom:1rem;
    display:flex; align-items:center; justify-content:space-between; gap:12px;
}
/* Icon badge for each block title — matching SatiCast's colored section-badge */
.block-icon-badge {
    width:34px; height:34px; border-radius:10px; flex-shrink:0;
    display:inline-flex; align-items:center; justify-content:center;
    font-size:1rem; margin-right:0.6rem; transition:transform 0.25s ease;
}
.colored-block:hover .block-icon-badge { transform:rotate(-8deg) scale(1.1); }
.narrative-block .block-icon-badge { background:linear-gradient(135deg,#E0E7FF,#C7D2FE); }
.chat-block .block-icon-badge      { background:linear-gradient(135deg,#FCE7F3,#FBCFE8); }
.whatsapp-block .block-icon-badge  { background:linear-gradient(135deg,#D1FAE5,#A7F3D0); }
.email-block .block-icon-badge     { background:linear-gradient(135deg,#FEF3C7,#FDE68A); }
.tomorrow-block .block-icon-badge  { background:linear-gradient(135deg,#CFFAFE,#A5F3FC); }
.blocker-block .block-icon-badge   { background:linear-gradient(135deg,#FEE2E2,#FECACA); }
.block-title-text { display:flex; align-items:center; }

.narrative-title { color:#4f46e5; } .chat-title { color:#db2777; }
.whatsapp-title  { color:#128C7E; } .email-title { color:#d97706; }
.tomorrow-title  { color:#0891b2; } .blocker-title { color:#dc2626; }

.action-btn {
    display:inline-flex; align-items:center; gap:6px;
    background:rgba(0,0,0,0.06); border:1.5px solid rgba(0,0,0,0.1); border-radius:999px;
    padding:0.4rem 1.1rem; font-size:0.8rem; font-weight:700;
    color:var(--text-h); text-decoration:none; transition:all 0.2s;
}
.action-btn:hover { background:rgba(217,119,87,0.1); border-color:rgba(217,119,87,0.3); transform:translateY(-1px); }

/* ── CODE BLOCKS ── */
pre {
    border-radius:16px !important; background:var(--code-bg) !important;
    border:1.5px solid rgba(0,0,0,0.06) !important; padding:1.25rem !important;
    color:var(--code-txt) !important; font-size:0.94rem !important; line-height:1.75 !important;
}

/* ── EDITABLE OUTPUT TEXTAREAS (narrative/chat/whatsapp/email) — matches
   the previous read-only code-block look, but content is now editable. ── */
.colored-block .stTextArea textarea {
    background:var(--code-bg) !important; color:var(--code-txt) !important;
    border-radius:16px !important; border:1.5px solid rgba(0,0,0,0.06) !important;
    font-family:'SFMono-Regular',Consolas,monospace !important;
    font-size:0.9rem !important; line-height:1.7 !important;
}

/* ── CARD HOVER ACCENT SWEEP — matching claude.com/blog's card hover style ── */
.custom-card, .colored-block, .hist-card, .member-chip {
    position:relative; overflow:hidden;
}
.custom-card::before, .colored-block::before {
    content:'';
    position:absolute; top:0; left:0; height:3px; width:0;
    background:linear-gradient(90deg,#D97757,#B45532);
    transition:width 0.35s ease;
}
.custom-card:hover::before, .colored-block:hover::before { width:100%; }

/* ── ACCESSIBILITY: visible focus rings for keyboard navigation ── */
button:focus-visible, input:focus-visible, textarea:focus-visible, a:focus-visible {
    outline: 3px solid #D97757 !important;
    outline-offset: 2px !important;
}

/* ── CUSTOM SCROLLBAR ── */
::-webkit-scrollbar { width:10px; height:10px; }
::-webkit-scrollbar-track { background:transparent; }
::-webkit-scrollbar-thumb { background:rgba(217,119,87,0.35); border-radius:99px; }
::-webkit-scrollbar-thumb:hover { background:rgba(217,119,87,0.55); }

/* ── PRINT-FRIENDLY ── */
@media print {
    .stButton, #dmchk, .dm-label { display:none !important; }
    .stApp { background:#fff !important; }
    .colored-block, .custom-card { box-shadow:none !important; }
}

/* ── STATS ROW ── */
.stats-row { display:grid; grid-template-columns:repeat(3,1fr); gap:14px; margin:1.6rem 0 0; }
@media (max-width:600px) { .stats-row { grid-template-columns:1fr; } }
.stat-chip {
    background:var(--chip-bg); border:1.5px solid var(--chip-bdr); border-radius:16px;
    padding:0.9rem 1.2rem; backdrop-filter:blur(8px); text-align:center;
    transition:transform 0.2s ease;
}
.stat-chip:hover { transform:translateY(-3px); }
.stat-val {
    font-family:'Syne',sans-serif; font-size:1.7rem; font-weight:800; color:var(--text-h);
    display:inline-block; animation:statPopIn 0.5s cubic-bezier(0.34,1.56,0.64,1) both;
}
@keyframes statPopIn {
    0%   { opacity:0; transform:scale(0.4) translateY(8px); }
    60%  { opacity:1; transform:scale(1.15) translateY(-2px); }
    100% { opacity:1; transform:scale(1) translateY(0); }
}
.stat-chip:nth-child(1) .stat-val { animation-delay:0.05s; }
.stat-chip:nth-child(2) .stat-val { animation-delay:0.15s; }
.stat-chip:nth-child(3) .stat-val { animation-delay:0.25s; }
.stat-lbl { font-size:0.74rem; font-weight:700; letter-spacing:0.08em; text-transform:uppercase; color:var(--text-m); margin-top:0.15rem; }

/* ── TOKEN BADGE ── */
.token-row { text-align:center; margin:1.8rem 0 0; }
.token-badge {
    display:inline-block; background:rgba(255,255,255,0.18); border:1.5px solid rgba(255,255,255,0.35);
    border-radius:999px; padding:0.5rem 1.5rem; color:#fff !important; font-weight:700; font-size:0.86rem;
    backdrop-filter:blur(10px);
}

/* ── HISTORY ── */
.hist-card {
    background:var(--card-bg); border:1.5px solid var(--card-bdr); border-radius:18px;
    padding:1.1rem 1.4rem; margin-bottom:0.85rem; backdrop-filter:blur(12px);
    display:flex; justify-content:space-between; align-items:flex-start; gap:12px;
    transition:transform 0.2s ease;
}
.hist-card:hover { transform:translateY(-2px); }
.hist-date { font-family:'Syne',sans-serif; font-size:0.95rem; font-weight:800; color:var(--text-h); }
.hist-meta { font-size:0.76rem; color:var(--text-m); margin-top:0.2rem; }

/* ── EMPTY STATES — a consistent, illustrated placeholder instead of a
   blank card when there's nothing to show yet. ── */
.empty-state {
    text-align:center; padding:2.5rem 1.5rem; border-radius:18px;
    border:1.5px dashed var(--card-bdr); margin-top:1rem;
    animation:fadeUp 0.5s ease both;
}
.empty-state-icon { font-size:2.4rem; margin-bottom:0.6rem; opacity:0.7; }
.empty-state-title { font-weight:800; font-size:1rem; color:var(--text-h); margin-bottom:0.4rem; }
.empty-state-sub { font-size:0.82rem; color:var(--text-m); max-width:420px; margin:0 auto; line-height:1.5; }

/* ── VISUAL TIMELINE — dots + connecting line, an at-a-glance overview
   of saved standups before drilling into any single entry. ── */
.timeline-wrap {
    position:relative; display:flex; overflow-x:auto; gap:0;
    padding:1.6rem 0.5rem 0.8rem; margin-bottom:0.5rem;
}
.timeline-line {
    position:absolute; top:1.9rem; left:2rem; right:2rem; height:2px;
    background:linear-gradient(90deg, rgba(217,119,87,0.15), rgba(217,119,87,0.5), rgba(217,119,87,0.15));
}
.tl-node { position:relative; flex:0 0 92px; text-align:center; z-index:1; }
.tl-dot {
    width:14px; height:14px; border-radius:50%; background:#D97757; margin:0 auto 0.5rem;
    border:3px solid var(--card-bg); box-shadow:0 0 0 2px rgba(217,119,87,0.35);
    transition:transform 0.15s ease;
}
.tl-node:hover .tl-dot { transform:scale(1.35); }
.tl-label { font-size:0.68rem; font-weight:700; color:var(--text-h); white-space:nowrap; }
.tl-sub { font-size:0.62rem; color:var(--text-m); }

details { background:var(--card-bg) !important; border-radius:18px !important; border:1.5px solid var(--card-bdr) !important; padding:0.5rem 1rem !important; margin-bottom:1rem !important; }
details summary { color:var(--text-h) !important; font-weight:700 !important; cursor:pointer; }

@media (max-width:768px) {
    .main-title  { font-size:2.8rem; }
    .title-text  { font-size:1.2rem; }
    .side-emoji  { display:none; }
    .card-header { flex-direction:column; align-items:flex-start; gap:12px; }
    .loader-wrap { padding:2rem 1.25rem; }
}
</style>
"""

st.markdown(CSS, unsafe_allow_html=True)
st.markdown(
    '<input type="checkbox" id="dmchk" aria-label="Toggle dark mode">'
    '<label for="dmchk" class="dm-label" role="switch" aria-label="Dark mode switch" tabindex="0" '
    'onkeydown="if(event.key===\'Enter\'||event.key===\' \'){event.preventDefault();document.getElementById(\'dmchk\').click();}"></label>',
    unsafe_allow_html=True
)

# Command palette (Cmd/Ctrl+K) — same mechanism as SatiCast's: finds real
# elements in the parent document and clicks them, best-effort.
components.html("""
<script>
try {
  const doc = window.parent.document;
  if (!doc.getElementById('sanghaCmdPalette')) {
    const overlay = doc.createElement('div');
    overlay.id = 'sanghaCmdPalette';
    overlay.style.cssText = 'display:none;position:fixed;inset:0;z-index:999998;'
      + 'background:rgba(0,0,0,0.45);align-items:flex-start;justify-content:center;padding-top:12vh;';
    overlay.innerHTML = `
      <div style="background:var(--card-bg,#fff);border-radius:16px;padding:0.6rem;width:min(420px,90vw);
                  box-shadow:0 24px 60px rgba(0,0,0,0.3);font-family:'Inter',sans-serif;">
        <div style="padding:0.5rem 0.7rem;font-size:0.7rem;font-weight:700;letter-spacing:0.05em;
                    text-transform:uppercase;opacity:0.55;">Quick Actions &nbsp;·&nbsp; Esc to close</div>
        <button data-cmd="generate" class="sangha-cmd-item">✨ &nbsp;Generate Professional Status</button>
        <button data-cmd="darkmode" class="sangha-cmd-item">🌗 &nbsp;Toggle Dark Mode</button>
        <button data-cmd="top" class="sangha-cmd-item">⬆️ &nbsp;Scroll to Top</button>
      </div>`;
    doc.body.appendChild(overlay);
    const style = doc.createElement('style');
    style.textContent = '.sangha-cmd-item { display:block; width:100%; text-align:left; padding:0.7rem 0.8rem; '
      + 'border:none; background:transparent; border-radius:10px; cursor:pointer; font-size:0.9rem; '
      + 'color:inherit; margin-bottom:2px; } .sangha-cmd-item:hover { background:rgba(217,119,87,0.14); }';
    doc.head.appendChild(style);

    function closePalette() { overlay.style.display = 'none'; }
    function openPalette() { overlay.style.display = 'flex'; }
    overlay.addEventListener('click', (e) => { if (e.target === overlay) closePalette(); });

    overlay.querySelectorAll('.sangha-cmd-item').forEach(btn => {
      btn.addEventListener('click', () => {
        const cmd = btn.dataset.cmd;
        if (cmd === 'darkmode') {
          const cb = doc.getElementById('dmchk');
          if (cb) cb.click();
        } else if (cmd === 'top') {
          (doc.scrollingElement || doc.documentElement).scrollTo({ top: 0, behavior: 'smooth' });
        } else if (cmd === 'generate') {
          const btns = Array.from(doc.querySelectorAll('button'));
          const target = btns.find(b => b.textContent.includes('Generate Professional Status'));
          if (target) target.click();
        }
        closePalette();
      });
    });

    doc.defaultView.addEventListener('keydown', (e) => {
      if ((e.metaKey || e.ctrlKey) && e.key.toLowerCase() === 'k') {
        e.preventDefault();
        overlay.style.display === 'flex' ? closePalette() : openPalette();
      } else if (e.key === 'Escape') {
        closePalette();
      }
    });
  }
} catch (e) {}
</script>
""", height=0, width=0)

# Auto-detect the OS/browser color-scheme preference on first visit only.
# Once the person has clicked the toggle themselves, their explicit choice
# (saved in the parent page's localStorage) always wins over the OS setting.
components.html("""
<script>
try {
  const doc = window.parent.document;
  const cb  = doc.getElementById('dmchk');
  if (cb) {
    const KEY = 'sanghastatus_dm_user_set';
    const userSet = doc.defaultView.localStorage.getItem(KEY);
    if (!userSet) {
      const prefersDark = doc.defaultView.matchMedia
        && doc.defaultView.matchMedia('(prefers-color-scheme: dark)').matches;
      if (prefersDark && !cb.checked) { cb.click(); }
    }
    if (!cb.dataset.sanghaListenerBound) {
      cb.dataset.sanghaListenerBound = "1";
      cb.addEventListener('change', () => {
        doc.defaultView.localStorage.setItem(KEY, '1');
      });
    }
  }
} catch (e) {}
</script>
""", height=0, width=0)

st.markdown(
    '<div class="fs-toggle">'
    '<input type="radio" name="fontsize" id="fsSmall">'
    '<input type="radio" name="fontsize" id="fsNormal" checked>'
    '<input type="radio" name="fontsize" id="fsLarge">'
    '<label for="fsSmall" class="fs-btn">A⁻</label>'
    '<label for="fsNormal" class="fs-btn">A</label>'
    '<label for="fsLarge" class="fs-btn">A⁺</label>'
    '</div>',
    unsafe_allow_html=True
)

# ═══════════════════════════════════════════════════
# HERO
# ═══════════════════════════════════════════════════
st.markdown(
    '<div class="hero-section">'
    '<span class="hero-icon">🏛️</span>'
    '<div class="main-title">Sangha<span></span>Status</div>'
    '<div class="main-subtitle">Professional standup narratives, chat updates, WhatsApp messages &amp; emails — generated in seconds.</div>'
    '<div class="sangha-meaning">✦ Sangha — the Pali word for "community" or "assembly" &nbsp;·&nbsp; Focusing on the Team ✦</div>'
    '</div>',
    unsafe_allow_html=True
)

# ═══════════════════════════════════════════════════
# CONFIGURATION CARD
# ═══════════════════════════════════════════════════
st.markdown(
    '<div class="custom-card">'
    '<div class="card-header"><div class="card-title">'
    '<div class="icon-circle purple-icon">⚙️</div>'
    '<div><div class="title-text">Configuration</div>'
    '<div class="desc-text">Set the date, tone, domain, and output language for this standup.</div></div>'
    '</div><div class="side-emoji">🎛️</div></div>',
    unsafe_allow_html=True
)

cfg1, cfg2, cfg3, cfg4 = st.columns(4)
with cfg1:
    standup_date = st.date_input("📅 Standup Date", value=date.today(), key="pref_date")
with cfg2:
    tone_choice = st.selectbox("🎨 Tone", options=list(TONE_OPTIONS.keys()), index=1, key="pref_tone")
with cfg3:
    domain_choice = st.selectbox("🏢 Domain", options=list(DOMAIN_OPTIONS.keys()), index=0, key="pref_domain")
    _prev_raw = st.session_state.get("raw_input", "")
    _suggested_domain = suggest_domain(_prev_raw)
    if _suggested_domain and _suggested_domain != domain_choice:
        if st.button(f"💡 Use suggested: {_suggested_domain}", key="apply_suggested_domain"):
            st.session_state.pref_domain = _suggested_domain
            st.rerun()
with cfg4:
    lang_choice = st.selectbox("🌐 Output Language", options=list(OUTPUT_LANG_OPTIONS.keys()), index=0, key="pref_lang")

cfg5, cfg6, cfg7 = st.columns(3)
with cfg5:
    # Pre-fill from the most recent saved entry — a lightweight "remembered"
    # default rather than a separate storage mechanism, since the same
    # persistent history file already carries this information.
    _remembered_project = ""
    if st.session_state.history and st.session_state.history[0].get("project") not in (None, "—"):
        _remembered_project = st.session_state.history[0]["project"]
    project_name = st.text_input("📁 Project / Sprint Tag", value=_remembered_project,
                                 placeholder="e.g. Phoenix · Sprint 14", key="pref_project")
with cfg6:
    include_tomorrow = st.checkbox("📅 Include Tomorrow's Plan", value=False, key="pref_tomorrow")
with cfg7:
    include_blockers = st.checkbox("🚧 Highlight Blockers", value=True, key="pref_blockers")

if "pref_your_name" not in st.session_state:
    st.session_state.pref_your_name = load_shared_settings().get("your_name", "")
your_name = st.text_input("👤 Your Name (optional — used as 'Prepared by' in exports · shared with SatiCast)",
                          placeholder="e.g. Pranay", key="pref_your_name")
if your_name.strip():
    save_shared_settings({"your_name": your_name.strip()})

# ── SETTINGS EXPORT / IMPORT — same idea as SatiCast's: download your
# preferences as JSON, restore them on a fresh browser/session. ──
with st.expander("⚙️ Export / Import Settings", expanded=False):
    _settings_snapshot = {
        "pref_tone": st.session_state.get("pref_tone"),
        "pref_domain": st.session_state.get("pref_domain"),
        "pref_lang": st.session_state.get("pref_lang"),
        "pref_project": st.session_state.get("pref_project"),
        "pref_tomorrow": st.session_state.get("pref_tomorrow"),
        "pref_blockers": st.session_state.get("pref_blockers"),
        "pref_your_name": st.session_state.get("pref_your_name"),
    }
    ecol1, ecol2 = st.columns(2)
    with ecol1:
        st.download_button(
            "⬇️ Download my settings", data=json.dumps(_settings_snapshot, indent=2),
            file_name="sanghastatus_settings.json", mime="application/json",
            key="dl_settings_btn", use_container_width=True
        )
    with ecol2:
        _uploaded_settings = st.file_uploader("Restore from file", type=["json"], key="settings_upload", label_visibility="collapsed")
        if _uploaded_settings is not None:
            try:
                _restored = json.load(_uploaded_settings)
                for _k, _v in _restored.items():
                    if _v is not None:
                        st.session_state[_k] = _v
                st.success("✅ Settings restored — refresh above widgets by re-running.")
                st.rerun()
            except Exception as e:
                st.error(f"Couldn't read that settings file: {e}")

st.markdown("</div>", unsafe_allow_html=True)

# ═══════════════════════════════════════════════════
# RAW UPDATES CARD
# ═══════════════════════════════════════════════════
st.markdown(
    '<div class="custom-card">'
    '<div class="card-header"><div class="card-title">'
    '<div class="icon-circle green-icon">📝</div>'
    '<div><div class="title-text">Raw Updates</div>'
    '<div class="desc-text">Paste your team\'s raw notes below — works for engineering, functional, or HR teams. Names are auto-detected.</div></div>'
    '</div><div class="side-emoji">📋</div></div>',
    unsafe_allow_html=True
)

raw_updates = st.text_area(
    "Paste team updates below:",
    placeholder="""Alice:
- Working on API integration with payment gateway.
- Resolving backend timeout issues.
- Tomorrow: complete unit tests.

Bob:
- Preparing regression test cases for release.
- Blocked on UAT environment access from DevOps.

Carol:
- Deployed hotfix to staging.
- Reviewing PR for auth module.
""",
    label_visibility="collapsed",
    key="raw_input",
    help="🎙️ Tip: use your keyboard's dictation button (mic icon on the on-screen/OS keyboard) to speak your update — it types straight into this box like any other text."
)

if GITHUB_PULL_ENABLED:
    if st.button(f"🔽 Pull my GitHub activity (last 24h) — {GITHUB_USERNAME}", key="pull_github_btn"):
        gh_bullets = fetch_github_activity(24)
        if gh_bullets:
            existing = st.session_state.get("raw_input", "")
            block = f"{GITHUB_USERNAME}:\n" + "\n".join(gh_bullets)
            st.session_state.raw_input = (existing.rstrip() + "\n\n" + block).strip() if existing.strip() else block
            st.rerun()
        else:
            st.caption("No GitHub activity found in the last 24h, or the pull failed.")

_day_segments = split_multi_day_paste(raw_updates) if raw_updates.strip() else []
if _day_segments:
    st.caption(f"📚 Looks like {len(_day_segments)} days' worth of updates pasted at once — pick one to generate now (paste the rest again later, or use ➡️ Load to swap in a different day).")
    day_labels = [lbl for lbl, _ in _day_segments]
    dcol1, dcol2 = st.columns([3, 1])
    with dcol1:
        picked_day_label = st.selectbox("Which day to generate for right now?", options=day_labels, key="multi_day_pick")
    with dcol2:
        st.write("")
        if st.button("➡️ Load this day", key="load_multi_day", use_container_width=True):
            picked_raw = next(txt for lbl, txt in _day_segments if lbl == picked_day_label)
            st.session_state.raw_input = picked_raw
            st.rerun()

members = parse_members(raw_updates) if raw_updates.strip() else []

_warn = looks_malformed(raw_updates, members)
if _warn and raw_updates.strip():
    st.caption(f"⚠️ {_warn}")

if members:
    total_tasks    = sum(m["count"] for m in members)
    total_blockers = sum(len(m["blockers"]) for m in members)
    st.markdown(
        f'<div class="stats-row">'
        f'<div class="stat-chip"><div class="stat-val">{len(members)}</div><div class="stat-lbl">Members</div></div>'
        f'<div class="stat-chip"><div class="stat-val">{total_tasks}</div><div class="stat-lbl">Total Tasks</div></div>'
        f'<div class="stat-chip"><div class="stat-val">{total_blockers}</div><div class="stat-lbl">Blockers</div></div>'
        f'</div>',
        unsafe_allow_html=True
    )
    _roster = load_roster()
    chips = ""
    avatar_colors = ["#D97757", "#B45532", "#23a6d5", "#23d5ab", "#f59e0b", "#ef4444"]
    # Role → accent color, reusing the same keyword families as the domain
    # detector so a "Backend Engineer" and a "QA Lead" get visually
    # distinct stripes even within one mixed team.
    ROLE_ACCENT_PALETTE = {
        "dev": "#23a6d5", "engineer": "#23a6d5", "backend": "#23a6d5", "frontend": "#23a6d5",
        "qa": "#23d5ab", "test": "#23d5ab",
        "devops": "#f59e0b", "infra": "#f59e0b", "sre": "#f59e0b",
        "design": "#ec4899", "ux": "#ec4899",
        "product": "#8b5cf6", "pm": "#8b5cf6",
        "hr": "#14b8a6", "recruit": "#14b8a6",
        "finance": "#84cc16",
    }

    def _role_accent(role: str) -> str:
        role_l = role.lower()
        for kw, color in ROLE_ACCENT_PALETTE.items():
            if kw in role_l:
                return color
        return "transparent"

    for i, m in enumerate(members):
        bl = (f'<div class="chip-block">⚠️ {len(m["blockers"])} blocker{"s" if len(m["blockers"])>1 else ""}</div>'
              if m["blockers"] else "")
        initial = m["name"].strip()[:1].upper() or "?"
        color = avatar_colors[i % len(avatar_colors)]
        role = _roster.get(m["name"].lower(), {}).get("role", "")
        role_html = f'<div class="chip-tasks" style="opacity:0.75;">{role}</div>' if role else ""
        accent = _role_accent(role)
        chips += (
            f'<div class="member-chip" style="display:flex;align-items:center;gap:10px;'
            f'border-left:4px solid {accent};">'
            f'<div class="chip-count-badge">{m["count"]}</div>'
            f'<div style="width:32px;height:32px;border-radius:50%;background:{color};'
            f'color:#fff;display:flex;align-items:center;justify-content:center;'
            f'font-weight:800;font-size:0.9rem;flex-shrink:0;">{initial}</div>'
            f'<div><div class="chip-name">{m["name"]}</div>{role_html}'
            f'<div class="chip-tasks">{m["count"]} task{"s" if m["count"]!=1 else ""}</div>{bl}</div></div>'
        )
    st.markdown(f'<div class="member-chips">{chips}</div>', unsafe_allow_html=True)

    # ── TEAM ROSTER — profile memory (role tag per member, remembered
    # across runs so it doesn't need re-entering every standup) ──
    with st.expander("👥 Team Roster — tag each member's role (remembered next time)", expanded=False):
        roster_changed = False
        for m in members:
            key = m["name"].lower()
            current_role = _roster.get(key, {}).get("role", "")
            new_role = st.text_input(
                f"Role for {m['name']}", value=current_role,
                placeholder="e.g. Backend Engineer, QA Lead…",
                key=f"roster_role_{key}"
            )
            if new_role.strip() != current_role:
                _roster[key] = {"display": m["name"], "role": new_role.strip()}
                roster_changed = True
        if roster_changed:
            save_roster(_roster)

    # ── DIFF AGAINST YESTERDAY — actual side-by-side before/after ──
    if st.session_state.history:
        diffs = diff_against_previous(members, st.session_state.history[0])
        if diffs:
            rows = ""
            for name, pairs in diffs.items():
                for today_task, yesterday_task in pairs:
                    old_html, new_html = word_diff_html(yesterday_task, today_task)
                    rows += (
                        f'<div class="diff-row">'
                        f'<div class="diff-name">{name}</div>'
                        f'<div class="diff-cols">'
                        f'<div class="diff-col diff-yesterday"><span class="diff-label">Yesterday</span>{old_html}</div>'
                        f'<div class="diff-arrow">→</div>'
                        f'<div class="diff-col diff-today"><span class="diff-label">Today</span>{new_html}</div>'
                        f'</div></div>'
                    )
            st.markdown(
                f'<div class="diff-panel">'
                f'<div class="diff-panel-title">🔁 Possibly carried over from last update</div>'
                f'{rows}</div>',
                unsafe_allow_html=True
            )

        # ── BRIEF-UPDATE NOTICE — purely a word-count comparison against the
        # person's own history, NOT a claim about mood or wellbeing. ──
        brief_flags = flag_brief_updates(members, st.session_state.history)
        if brief_flags:
            brief_lines = "".join(
                f'<div style="margin-top:0.3rem;"><strong>{name}:</strong> '
                f'~{today}w vs usual ~{usual}w</div>'
                for name, (today, usual) in brief_flags.items()
            )
            st.markdown(
                f'<div style="margin-top:0.6rem;background:rgba(59,130,246,0.06);'
                f'border:1.5px solid rgba(59,130,246,0.2);border-radius:12px;'
                f'padding:0.7rem 1rem;font-size:0.85rem;color:#1E3A8A;">'
                f'📏 <strong>Notably briefer than usual (word count only):</strong>{brief_lines}</div>',
                unsafe_allow_html=True
            )
else:
    st.markdown(
        '<div class="empty-state">'
        '<div class="empty-state-icon">📝</div>'
        '<div class="empty-state-title">Paste your team\'s updates to get started</div>'
        '<div class="empty-state-sub">One line per person as a header (e.g. "Alice:"), followed by bullet points — '
        'members, tasks and blockers are detected automatically.</div>'
        '</div>',
        unsafe_allow_html=True
    )

st.markdown("</div>", unsafe_allow_html=True)

# ═══════════════════════════════════════════════════
# GENERATE
# ═══════════════════════════════════════════════════
generate = st.button("✨ Generate Professional Status")

if generate:
    if not raw_updates.strip():
        st.warning("⚠️ Please paste your team's raw updates first.")
    elif not members:
        st.warning("⚠️ Could not detect any team member names. Use 'Name:' format.")
    else:
        slot = st.empty()
        slot.markdown(render_loader(0), unsafe_allow_html=True)

        try:
            fmt_date    = standup_date.strftime("%B %d, %Y")
            project_tag = f" | {project_name}" if project_name.strip() else ""
            tone_instr  = TONE_OPTIONS[tone_choice]
            domain_str  = DOMAIN_OPTIONS[domain_choice]
            lang_str    = OUTPUT_LANG_OPTIONS[lang_choice]

            prompt_payload = f"""
Use this date for all sections: {fmt_date}
Project/Sprint: {project_name or "General"}
Team Members detected: {len(members)} ({', '.join(m['name'] for m in members)})

Raw Updates:
{raw_updates}

Embed the project tag "{project_tag}" in the chat_update header and email subject line.
"""
            narrative_prompt = build_narrative_prompt(tone_instr, domain_str, lang_str, include_tomorrow, include_blockers)
            chat_prompt       = build_chat_prompt(tone_instr, domain_str, lang_str)
            email_prompt      = build_email_prompt(tone_instr, domain_str, lang_str)

            slot.markdown(render_loader(1), unsafe_allow_html=True)

            # Scale token budgets by team size — a 1-person update needs far
            # fewer tokens than a 10-person one; fixed budgets either waste
            # headroom (small teams) or truncate (large teams).
            team_size = max(1, len(members))
            narrative_tokens = min(3000, 900 + team_size * 220)
            chat_tokens       = min(2600, 800 + team_size * 180)
            email_tokens      = min(2200, 700 + team_size * 160)

            # ── THREE INDEPENDENT CALLS IN PARALLEL ──
            # Narrative, chat/WhatsApp, and email don't depend on each other,
            # so firing them concurrently cuts total wait time to roughly the
            # slowest single call instead of the sum of all three.
            total_tokens_used = 0
            with ThreadPoolExecutor(max_workers=3) as ex:
                fut_narrative = ex.submit(call_llm_json, narrative_prompt, prompt_payload, 0.1, narrative_tokens)
                fut_chat      = ex.submit(call_llm_json, chat_prompt, prompt_payload, 0.1, chat_tokens)
                fut_email     = ex.submit(call_llm_json, email_prompt, prompt_payload, 0.1, email_tokens)

                results = {}
                errors  = {}
                for name, fut in (("narrative", fut_narrative), ("chat", fut_chat), ("email", fut_email)):
                    try:
                        piece, completion = fut.result()
                        results[name] = piece
                        if hasattr(completion, "usage") and completion.usage:
                            total_tokens_used += completion.usage.total_tokens
                    except Exception as e:
                        errors[name] = e

                # Retry only the piece(s) that actually failed, not all three —
                # cheaper and faster than a blanket full retry.
                retry_prompts = {
                    "narrative": (narrative_prompt, narrative_tokens),
                    "chat":      (chat_prompt, chat_tokens),
                    "email":     (email_prompt, email_tokens),
                }
                for name in list(errors.keys()):
                    p, budget = retry_prompts[name]
                    try:
                        piece, completion = call_llm_json(p, prompt_payload, 0.05, budget)
                        results[name] = piece
                        if hasattr(completion, "usage") and completion.usage:
                            total_tokens_used += completion.usage.total_tokens
                        del errors[name]
                    except Exception:
                        pass  # still failed after retry — that piece just won't render

                data = {}
                for piece in results.values():
                    data.update(piece)

                if errors:
                    st.warning(f"⚠️ Could not generate: {', '.join(errors.keys())}. Try regenerating.")

            slot.markdown(render_loader(2), unsafe_allow_html=True)

            st.session_state.session_tokens += total_tokens_used

            # Capture yesterday's entry (if any) BEFORE inserting today's,
            # so we can diff today's tasks against it.
            previous_entry = st.session_state.history[0] if st.session_state.history else None

            st.session_state.history.insert(0, {
                "date": fmt_date, "project": project_name or "—",
                "members": len(members), "tasks": sum(m["count"] for m in members),
                "tone": tone_choice, "domain": domain_choice, "lang": lang_choice,
                "data": data, "members_detail": members,
            })
            st.session_state.history = st.session_state.history[:MAX_HISTORY]
            save_history_to_disk(st.session_state.history)

            slot.empty()
            st.markdown("<br>", unsafe_allow_html=True)

            narrative = data.get("standup_narrative", "")
            chat      = data.get("chat_update", "")
            wa        = data.get("whatsapp_update", "")
            email_raw = data.get("email_update", "")

            lines      = email_raw.splitlines()
            subj_line  = next((l for l in lines if l.lower().startswith("subject:")), "")
            subj       = subj_line.replace("Subject:", "").replace("subject:", "").strip()
            body_lines = [l for l in lines if l != subj_line]
            body_text  = "\n".join(body_lines).strip()
            mailto     = mailto_link(subj, body_text)

            st.markdown('<div class="output-grid">', unsafe_allow_html=True)

            # Narrative — full width
            st.markdown(
                '<div class="colored-block narrative-block">'
                '<div class="block-title narrative-title"><span class="block-title-text"><span class="block-icon-badge">🗣️</span>Standup Narrative</span></div>',
                unsafe_allow_html=True
            )
            st.text_area("Narrative", narrative, height=dynamic_ta_height(narrative), key="edit_narrative", label_visibility="collapsed")
            render_copy_button(narrative, "narrative")
            st.markdown("</div>", unsafe_allow_html=True)

            # Chat + WhatsApp side by side
            st.markdown(
                '<div class="colored-block chat-block">'
                '<div class="block-title chat-title"><span class="block-title-text"><span class="block-icon-badge">💬</span>Chat Update (Slack / Teams)</span></div>',
                unsafe_allow_html=True
            )
            st.text_area("Chat", chat, height=dynamic_ta_height(chat), key="edit_chat", label_visibility="collapsed")
            render_copy_button(chat, "chat")
            with st.expander("🔗 Send directly to Slack", expanded=False):
                st.caption(
                    "Paste your Slack **Incoming Webhook URL** (Slack → Apps → Incoming Webhooks) — "
                    "it's used once for this send and isn't stored anywhere."
                )
                slack_webhook_url = st.text_input(
                    "Slack Webhook URL", placeholder="https://hooks.slack.com/services/…",
                    key="slack_webhook_url", type="password", label_visibility="collapsed"
                )
                if st.button("📤 Send to Slack", key="send_slack_btn"):
                    if not slack_webhook_url.strip():
                        st.warning("Paste a webhook URL first.")
                    elif not slack_webhook_url.startswith("https://hooks.slack.com/"):
                        st.warning("That doesn't look like a Slack webhook URL (should start with https://hooks.slack.com/…).")
                    else:
                        try:
                            import requests as _requests
                            resp = _requests.post(
                                slack_webhook_url.strip(),
                                json={"text": st.session_state.get("edit_chat", chat)},
                                timeout=8
                            )
                            if resp.status_code == 200:
                                st.success("✅ Sent to Slack!")
                            else:
                                st.error(f"Slack responded with {resp.status_code}: {resp.text[:200]}")
                        except Exception as e:
                            st.error(f"Couldn't reach Slack: {e}")
            st.markdown("</div>", unsafe_allow_html=True)

            st.markdown(
                '<div class="colored-block whatsapp-block">'
                '<div class="block-title whatsapp-title"><span class="block-title-text"><span class="block-icon-badge">📱</span>WhatsApp Update</span></div>',
                unsafe_allow_html=True
            )
            st.text_area("WhatsApp", wa, height=dynamic_ta_height(wa), key="edit_whatsapp", label_visibility="collapsed")
            render_copy_button(wa, "whatsapp")
            st.markdown("</div>", unsafe_allow_html=True)

            # Jira/Confluence markup — derived instantly from chat_update, no extra API call
            jira_markup = to_jira_confluence_markup(chat)
            st.markdown(
                '<div class="colored-block tomorrow-block">'
                '<div class="block-title tomorrow-title"><span class="block-title-text"><span class="block-icon-badge">🧩</span>Jira / Confluence Markup</span></div>',
                unsafe_allow_html=True
            )
            st.text_area("Jira", jira_markup, height=dynamic_ta_height(jira_markup), key="edit_jira", label_visibility="collapsed")
            st.markdown("</div>", unsafe_allow_html=True)

            # Email — full width with mailto button
            st.markdown(
                f'<div class="colored-block email-block">'
                f'<div class="block-title email-title">'
                f'<span class="block-title-text"><span class="block-icon-badge">📧</span>Email Update</span>'
                f'<a class="action-btn" href="{mailto}">✉️ Open in Mail App</a></div>',
                unsafe_allow_html=True
            )
            st.text_area("Email", email_raw, height=dynamic_ta_height(email_raw), key="edit_email", label_visibility="collapsed")
            render_copy_button(email_raw, "email")
            st.markdown("</div>", unsafe_allow_html=True)

            if include_tomorrow and data.get("tomorrow_plan"):
                st.markdown(
                    '<div class="colored-block tomorrow-block">'
                    '<div class="block-title tomorrow-title"><span class="block-title-text"><span class="block-icon-badge">📅</span>Tomorrow\'s Plan</span></div>',
                    unsafe_allow_html=True
                )
                st.code(data["tomorrow_plan"], language="text")
                st.markdown("</div>", unsafe_allow_html=True)

            if include_blockers and data.get("blocker_summary"):
                st.markdown(
                    '<div class="colored-block blocker-block">'
                    '<div class="block-title blocker-title"><span class="block-title-text"><span class="block-icon-badge">🚧</span>Blockers Summary</span></div>',
                    unsafe_allow_html=True
                )
                st.code(data["blocker_summary"], language="text")
                st.markdown("</div>", unsafe_allow_html=True)

            st.markdown("</div>", unsafe_allow_html=True)  # close output-grid

            run_tokens = total_tokens_used
            st.markdown(
                f'<div class="token-row"><div class="token-badge">'
                f'🛡️ This run: {run_tokens:,} tokens &nbsp;·&nbsp; '
                f'Session total: {st.session_state.session_tokens:,} tokens &nbsp;·&nbsp; '
                f'Runs saved: {len(st.session_state.history)}</div></div>',
                unsafe_allow_html=True
            )

            docx_bytes, docx_err = build_docx_export(data, fmt_date, prepared_by=your_name.strip())
            if docx_bytes:
                st.download_button(
                    "⬇️ Download as Word (.docx)", data=docx_bytes,
                    file_name=f"status_{date.today().isoformat()}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    key="dl_docx"
                )
            elif docx_err:
                st.caption(f"📄 Word export unavailable: {docx_err}")

        except json.JSONDecodeError:
            slot.empty()
            st.warning("⚠️ Malformed response — retrying once…")
            try:
                with ThreadPoolExecutor(max_workers=3) as ex:
                    fut_n = ex.submit(call_llm_json, narrative_prompt, prompt_payload, 0.05, 1600)
                    fut_c = ex.submit(call_llm_json, chat_prompt, prompt_payload, 0.05, 1400)
                    fut_e = ex.submit(call_llm_json, email_prompt, prompt_payload, 0.05, 1200)
                    data2 = {}
                    for fut in (fut_n, fut_c, fut_e):
                        piece, _ = fut.result()
                        data2.update(piece)
                st.success("✅ Retry succeeded!")
                st.code(data2.get("standup_narrative", ""), language="text")
            except Exception:
                st.error("❌ Both attempts failed. Please try again.")
        except Exception as e:
            slot.empty()
            st.error(f"❌ Error: {e}")

# ═══════════════════════════════════════════════════
# HISTORY
# ═══════════════════════════════════════════════════
if st.session_state.history:
    st.markdown("<br>", unsafe_allow_html=True)

    # ── VISUAL TIMELINE — dots + connecting line, newest first, as an
    # at-a-glance overview before drilling into any single entry below. ──
    _tl_dots = "".join(
        f'<div class="tl-node">'
        f'<div class="tl-dot" title="{e["date"]}"></div>'
        f'<div class="tl-label">{e["date"].split(" · ")[0]}</div>'
        f'<div class="tl-sub">{e.get("tasks", 0)} tasks</div>'
        f'</div>'
        for e in st.session_state.history
    )
    st.markdown(
        f'<div class="timeline-wrap"><div class="timeline-line"></div>{_tl_dots}</div>',
        unsafe_allow_html=True
    )

    with st.expander(f"🕘 Past Standups ({len(st.session_state.history)} saved) — click to expand & replay", expanded=False):
        for idx, entry in enumerate(st.session_state.history):
            st.markdown(
                f'<div class="hist-card"><div>'
                f'<div class="hist-date">📅 {entry["date"]} &nbsp;·&nbsp; {entry.get("project","—")}</div>'
                f'<div class="hist-meta">👥 {entry["members"]} members &nbsp;·&nbsp; ✅ {entry["tasks"]} tasks '
                f'&nbsp;·&nbsp; {entry["tone"]} &nbsp;·&nbsp; {entry["domain"]} &nbsp;·&nbsp; {entry["lang"]}</div>'
                f'</div></div>',
                unsafe_allow_html=True
            )
            with st.expander(f"View outputs — {entry['date']}", expanded=False):
                d = entry["data"]
                st.markdown("**🗣️ Narrative**")
                st.code(d.get("standup_narrative", ""), language="text")
                st.markdown("**💬 Chat**")
                st.code(d.get("chat_update", ""), language="text")
                st.markdown("**📱 WhatsApp**")
                st.code(d.get("whatsapp_update", ""), language="text")
                st.markdown("**📧 Email**")
                st.code(d.get("email_update", ""), language="text")
                if d.get("tomorrow_plan"):
                    st.markdown("**📅 Tomorrow**")
                    st.code(d["tomorrow_plan"], language="text")
                if d.get("blocker_summary"):
                    st.markdown("**🚧 Blockers**")
                    st.code(d["blocker_summary"], language="text")

        if st.button("🗑️ Clear all history", key="clear_hist"):
            st.session_state.history = []
            save_history_to_disk(st.session_state.history)
            st.rerun()

    # ── TEAM VELOCITY CHART — tasks reported per saved update, chronological ──
    if len(st.session_state.history) >= 2:
        import pandas as pd
        chron = list(reversed(st.session_state.history))  # oldest first for the chart
        chart_df = pd.DataFrame({
            "Date":  [e["date"] for e in chron],
            "Tasks": [e.get("tasks", 0) for e in chron],
        }).set_index("Date")
        st.markdown(
            '<div style="font-weight:700;font-size:0.85rem;letter-spacing:0.04em;'
            'text-transform:uppercase;color:var(--text-h);margin:1.5rem 0 0.5rem;">'
            '📈 Team Velocity — Tasks per Saved Update</div>',
            unsafe_allow_html=True
        )
        st.line_chart(chart_df, height=200)
        st.download_button(
            "⬇️ Download velocity data (.csv)",
            data=chart_df.reset_index().to_csv(index=False),
            file_name="team_velocity.csv", mime="text/csv",
            key="dl_velocity_csv"
        )

        # ── PER-PERSON TREND VIEW — expands the brevity flag into a full
        # per-member chart (task count + avg words/task across saved runs),
        # so a manager can see one person's pattern over time, not just
        # today's single-run flag.
        all_names = sorted({
            m["name"] for e in st.session_state.history
            for m in e.get("members_detail", [])
        })
        if all_names:
            with st.expander("👤 Per-Person Trend View", expanded=False):
                picked_member = st.selectbox("Team member", options=all_names, key="trend_member_pick")
                rows = []
                for e in chron:
                    match = next((m for m in e.get("members_detail", []) if m["name"] == picked_member), None)
                    if match:
                        n_tasks = len(match["tasks"])
                        avg_words = round(sum(len(t.split()) for t in match["tasks"]) / n_tasks, 1) if n_tasks else 0
                        rows.append({"Date": e["date"], "Tasks": n_tasks, "Avg words/task": avg_words})
                if len(rows) >= 2:
                    trend_df = pd.DataFrame(rows).set_index("Date")
                    st.line_chart(trend_df, height=200)
                    st.caption(f"📌 {picked_member} appears in {len(rows)} of {len(st.session_state.history)} saved updates.")
                    st.download_button(
                        f"⬇️ Download {picked_member}'s trend data (.csv)",
                        data=trend_df.reset_index().to_csv(index=False),
                        file_name=f"{picked_member.lower().replace(' ', '_')}_trend.csv", mime="text/csv",
                        key="dl_trend_csv"
                    )
                else:
                    st.caption(f"Not enough saved history for {picked_member} yet — need at least 2 updates mentioning them.")

    # ── WEEKLY ROLLUP — aggregates whatever's currently saved in history ──
    if len(st.session_state.history) >= 2:
        if st.button("📊 Generate Weekly Rollup", key="weekly_rollup_btn"):
            rollup_lines = ["WEEKLY ROLLUP", "=" * 40, ""]
            person_totals = {}
            for entry in st.session_state.history:
                rollup_lines.append(f"📅 {entry['date']} — {entry.get('project','—')}")
                narrative = entry["data"].get("standup_narrative", "")
                rollup_lines.append(narrative)
                rollup_lines.append("-" * 30)
                person_totals[entry.get("project", "—")] = person_totals.get(entry.get("project", "—"), 0) + entry.get("tasks", 0)
            rollup_lines.append("")
            rollup_lines.append(f"Total standups included: {len(st.session_state.history)}")
            rollup_lines.append(f"Total tasks across all entries: {sum(e.get('tasks',0) for e in st.session_state.history)}")
            rollup_text = "\n".join(rollup_lines)
            st.markdown(
                '<div class="colored-block narrative-block">'
                '<div class="block-title narrative-title"><span class="block-title-text"><span class="block-icon-badge">📊</span>Weekly Rollup</span></div>',
                unsafe_allow_html=True
            )
            st.code(rollup_text, language="text")
            st.download_button("⬇️ Download Rollup", data=rollup_text,
                              file_name=f"weekly_rollup_{date.today().isoformat()}.txt",
                              mime="text/plain", key="dl_rollup")
            st.markdown("</div>", unsafe_allow_html=True)

# ═══════════════════════════════════════════════════
# SUITE FOOTER — cross-link back to SatiCast. Only shown when the deployer
# has set SATICAST_URL in secrets.
# ═══════════════════════════════════════════════════
_sibling_url = st.secrets.get("SATICAST_URL", "")
if _sibling_url:
    st.markdown(
        f'<div style="text-align:center;margin:2.5rem 0 1rem;padding-top:1.2rem;'
        f'border-top:1px solid rgba(217,119,87,0.15);font-size:0.8rem;opacity:0.75;">'
        f'🏛️ SanghaStatus &nbsp;·&nbsp; part of the same suite as '
        f'<a href="{_sibling_url}" target="_blank" style="color:#D97757;font-weight:700;">🪷 SatiCast</a>'
        f'</div>',
        unsafe_allow_html=True
    )